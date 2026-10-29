#!/usr/bin/env python3
"""Presentation tables and charts: our Snake vs the Snake paper and SnakeCharmer
(docs/snake/results/slides/Snake_presentation_<date>.xlsx, one PNG per slide
in the same folder).

    PYTHONPATH=<dir with openpyxl> python3 util/snake/make_slides_data.py [<date>]

No simulations: every number is read by this script from existing outputs and
results files; make_report_xlsx.py adds the same sheets to the results
workbook.  Sheets:
  vs Snake paper                Tier 1 (PAPER_V100, 1B, Snake exact footprint):
                                sim_run_baseline, sim_run_snake, perfect L1D
                                from sim_run_snake_gate/D1_pl1d; paper values
                                from SPEC.md §5
  vs SnakeCharmer               SK match Step 1 (base QV100, 100M):
                                sim_run_snake_skm/qv100, no-prefetch runs from
                                sim_run_snake_skc; SnakeCharmer's SK averages
                                from results/SK_SC.md §4 (the eviction share
                                from the task brief: it is not in our docs)
  Per 100 prefetched sectors    derived from the previous sheet
  Best Snake vs no prefetching  SK_R0_D0 per workload
  Footnotes, Provenance         definitions, setup, caveats; every value's
                                source, build tags, run dirs and doc check
Each value is cross-checked against the results docs (final_tier1.md §3,
final_tier1.csv, D1_ceilings.md §5, SK_SC.md §4, SK_match_qv100.md).  A
missing value is "n/a" with the reason in Notes.  The saved workbook is
reopened: row counts and five values are checked against their sources.
Exit status 1 on a mismatch or a failed check.
"""
import csv
import datetime
import glob
import io
import os
import re
import subprocess
import sys
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import make_report_xlsx as R  # noqa: E402  (sheet format, Tier 1 parsers)
from s4_report import parse  # noqa: E402
from sk_sc_report import metrics, build, short, SEVEN, INPUT_DIFFERS, MISDECODED  # noqa: E402
from sk_match_report import CONFIGS as SKM_CONFIGS, Q  # noqa: E402

import matplotlib  # noqa: E402
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.drawing.image import Image as XLImage  # noqa: E402

ROOT, DOCS = R.ROOT, R.DOCS
RES = os.path.join(DOCS, "results")
PL1 = R.B + "-PERFECT_L1D"
D1_PL1 = os.path.join("sim_run_snake_gate", "D1_pl1d")
SKM = os.path.join("sim_run_snake_skm", "qv100")
SKC = "sim_run_snake_skc"
SK_LABELS = [("SK", "SK (reserve 8, decoupling on)"), ("SK_R0", "SK_R0 (reserve 0)"),
             ("SK_R0_D0", "SK_R0_D0 (reserve 0, decoupling off)")]
BEST = "SK_R0_D0"
NAIVE = "particlefilter_naive-rodinia-3.1"
VER, PRE, DER, PUB = R.VER, "Preliminary (100M run)", "Derived", "Published value"
R.FILL.update({PRE: R.FILL[R.PRE], DER: "DDEBF7", PUB: "E4DFEC"})
R.FMT["dec1"] = "0.0"
NAMES = {"s1": "vs Snake paper", "s2": "vs SnakeCharmer", "s3": "Per 100 prefetched sectors",
         "s4": "Best Snake vs no prefetching", "foot": "Footnotes", "prov": "Provenance"}
REPORT_NAMES = dict(NAMES, foot="Slides_footnotes", prov="Slides_provenance")
# SnakeCharmer's SK averages as listed in the task brief (2026-10-27, A100,
# 14 apps).  All but the eviction share are read from results/SK_SC.md §4 and
# only checked against these; the eviction share is in none of our docs.
BRIEF_SC = {"cov": 11.5, "acc": 23.1, "on": 48.3, "early": 28.7, "late": 23.0, "evict": 34.0}
CHECKS = []     # (where, ok)
MISMATCH = []   # (where, ours, doc)
USED = {}       # run dir -> {build tags: {job dirs}}


class V:
    """One table value: where it came from, how it was computed, the build
    tags of the runs behind it and its checks against the results docs."""

    def __init__(self, v, src, how="", tags=()):
        self.v, self.src, self.how, self.tags, self.chk = v, src, how, set(tags), []

    def check(self, doc, decimals, where):
        return check(self, self.v, doc, decimals, where)


def check(v, ours, doc, decimals, where):
    """ours, rounded like the doc, vs a doc value; logged on V v (or None)."""
    ok = ours is not None and doc is not None and \
        abs(round(ours, decimals) - doc) <= 0.5 * 10 ** -decimals + 1e-9
    if v is not None:
        v.chk.append("%s %s (%s)" % ("=" if ok else "MISMATCH: doc", doc, where))
    CHECKS.append((where, ok))
    if not ok:
        MISMATCH.append((where, ours, doc))
    return ok


def nan2none(x):
    return None if x is None or x != x else x


def rel(p):
    return os.path.relpath(p, ROOT)


def git(*args):
    return subprocess.run(["git", "-C", ROOT] + list(args), capture_output=True, text=True).stdout.strip()


def gm(gains):
    """GMEAN of percent gains (None if empty)."""
    g = R.gmean([1 + x / 100 for x in gains if x is not None])
    return (g - 1) * 100 if g else None


def use(path, tags):
    k = " ".join(sorted(t for t in tags if t))
    USED.setdefault(run_dir_of(path), {}).setdefault(k, set()).add(os.path.dirname(path) if
                                                                   path.endswith("sim.out") else path)


def run_dir_of(path):
    """Run dir (sim_run_*/...) of a job path: the part before /runs/."""
    r = rel(path)
    return r.split(os.sep + "runs" + os.sep)[0]


def opts(job_dir):
    """GPU setup and instruction cap of a job from its gpgpusim.config (last
    uncommented value of each option)."""
    t = R.rd(os.path.join(job_dir, "gpgpusim.config"))
    g = lambda k: (re.findall(r"(?m)^-%s\s+(\S+)" % k, t) or [None])[-1]
    clk = g("gpgpu_clock_domains")
    insn = int(g("gpgpu_max_insn"))
    cap = ("%dB" % (insn // 10 ** 9)) if insn % 10 ** 9 == 0 else ("%dM" % (insn // 10 ** 6))
    gpu = "%d SMs, %s, %s MHz" % (int(g("gpgpu_n_clusters")) * int(g("gpgpu_n_cores_per_cluster")),
                                  g("gpgpu_scheduler").upper(), clk.split(":")[0].split(".")[0] if clk else "n/a")
    return gpu, "%s instructions (-gpgpu_max_insn %d; whole run if shorter)" % (cap, insn)


def trace_arch(job_dirs):
    """Trace architectures ("-binary version" of each job's first kernel)."""
    vers = set()
    for d in job_dirs:
        p = os.path.join(d, "traces", "kernel-1.traceg")
        if not os.path.exists(p):
            continue
        with open(p, errors="replace") as f:
            for _, line in zip(range(40), f):
                m = re.match(r"-binary version = (\d+)", line)
                if m:
                    vers.add("sm_" + m.group(1))
    return ", ".join(sorted(vers)) or "n/a (no kernel-1.traceg found)"


# ------------------------------------------------------------ slide 1
def t1_key(cell, WL):
    """(app, args) of a results-doc row label like 'hotspot / 1024_2_2___d'."""
    c = cell.replace("**", "")
    for a, g in WL:
        ap = R.app_of(a)
        if " / " in c and c.startswith(ap + " / ") and g.startswith(c.split(" / ")[1][:10]):
            return (a, g)
    return None


def slide1(d):
    dirs = {}
    for dr, cfg in R.tier1_list():
        dirs.setdefault(cfg, dr)
    paper, prow, _ = R.paper_values()
    pavg = {nm: v for nm, v, _ in prow}.get("Avg / GMEAN")
    spec = R.rd(os.path.join(DOCS, "SPEC.md"))
    m = re.search(r"Baseline L1 hit rate (\d+)% → Snake (\d+)% → Isolated-Snake (\d+)%", spec)
    fig25 = tuple(float(x) for x in m.groups()) if m else (None, None, None)
    SP = "docs/snake/SPEC.md §5"
    WL = R.workloads()
    tags = lambda r: {r["accelsim_tag"], r["gpgpusim_tag"]} if r else set()
    src = lambda *rs: "; ".join(rel(r["path"]) for r in rs if r) or "n/a"
    rows, per_app, jobs, d1 = [], {}, {}, {}
    for a, g in WL:
        ap = R.app_of(a)
        b = R.job(os.path.join(ROOT, dirs[R.B], "runs"), a, g, R.B)
        s = R.job(os.path.join(ROOT, dirs[R.S], "runs"), a, g, R.S)
        p = R.job(os.path.join(ROOT, D1_PL1, "runs"), a, g, PL1)
        jobs[(R.B, a, g)], jobs[(R.S, a, g)] = b, s
        for r in (b, s, p):
            if r:
                use(r["path"], tags(r))
        rt = R.snake_rates(s) if s else dict(cov_pred=None, cov_paper=None, acc=None)
        pv = paper.get(ap)
        P = lambda i, fig: V(pv[i] if pv else None, "%s (%s, read off the bar chart, ±3 pp)" % (SP, fig),
                             "paper value for %s" % ap)
        notes = ["paper values per app, read off bar charts (±3 pp)"]
        if ap in ("hotspot", "lud"):
            notes.append("both inputs share the paper's app value")
        notes.append("paper L1 hit: n/a, Fig. 25 gives only the average over the paper's 11 apps")
        for r, what, dr in ((b, "baseline", dirs[R.B]), (s, "Snake", dirs[R.S]), (p, "perfect-L1D", D1_PL1)):
            if r is None:
                notes.append("%s run missing or unfinished in %s" % (what, dr))
        rows.append([R.wl_label(a, g), P(0, "Fig. 16"),
                     V(rt["cov_pred"], src(s), "snake_demand_pred_matched / snake_demand_sector_accesses", tags(s)),
                     V(rt["cov_paper"], src(s), "(snake_pf_useful_timely + snake_pf_late) / snake_demand_sector_accesses",
                       tags(s)),
                     P(1, "Fig. 17"),
                     V(rt["acc"], src(s), "snake_pf_useful_timely / snake_demand_sector_accesses", tags(s)),
                     P(2, "Fig. 18"),
                     V(R.gain(s, b), src(s, b), "gpu_tot_ipc Snake / gpu_tot_ipc baseline - 1", tags(s) | tags(b)),
                     V(R.gain(p, b), src(p, b), "gpu_tot_ipc perfect L1D / gpu_tot_ipc baseline - 1", tags(p) | tags(b)),
                     V(b and b["l1_hit_rate"], src(b), "GLOBAL_ACC_R HIT / TOTAL_ACCESS (baseline)", tags(b)),
                     V(s and s["l1_hit_rate"], src(s), "GLOBAL_ACC_R HIT / TOTAL_ACCESS (Snake)", tags(s)),
                     V(None, SP, "Fig. 25 gives the 11-app average only"),
                     V(None, SP, "Fig. 25 gives the 11-app average only"),
                     VER, "; ".join(notes)])
        r = rows[-1]
        per_app.setdefault(ap, []).append({k: r[i].v for k, i in (
            ("cov_pred", 2), ("cov_paper", 3), ("acc", 5), ("gain", 7), ("ceil", 8), ("l1b", 9), ("l1s", 10))})
    gma = lambda k: gm([gm([x[k] for x in xs]) for xs in per_app.values()])
    ama = lambda k: R.mean([R.mean([x[k] for x in xs]) for xs in per_app.values()])
    alltags = set().union(*[x.tags for row in rows for x in row if isinstance(x, V)])
    AM = "mean over the 7 apps (an app's inputs averaged first) of the rows above"
    GMH = "GMEAN over the 7 apps (an app's inputs combined first, geometric) of the rows above"
    PA = "%s; %s" % (SP, "mean over the same 7 apps")
    rows.append(["7 apps: GMEAN (IPC) / mean (rates)",
                 V(R.mean([paper[x][0] for x in R.APPS]), PA, "mean of the paper's Fig. 16 values"),
                 V(ama("cov_pred"), "this sheet", AM, alltags), V(ama("cov_paper"), "this sheet", AM, alltags),
                 V(R.mean([paper[x][1] for x in R.APPS]), PA, "mean of the paper's Fig. 17 values"),
                 V(ama("acc"), "this sheet", AM, alltags),
                 V(gm([paper[x][2] for x in R.APPS]), SP, "GMEAN over the same 7 apps of the paper's Fig. 18 values"),
                 V(gma("gain"), "this sheet", GMH, alltags), V(gma("ceil"), "this sheet", GMH, alltags),
                 V(ama("l1b"), "this sheet", AM, alltags), V(ama("l1s"), "this sheet", AM, alltags),
                 V(fig25[0], SP + " 'Other reference points' (§5.7, Fig. 25)", "paper average over its 11 apps"),
                 V(fig25[1], SP + " 'Other reference points' (§5.7, Fig. 25)", "paper average over its 11 apps"),
                 VER, "IPC: GMEAN over apps; rates: mean over apps. Paper columns: the same 7 apps (its own 11-app "
                 "averages: coverage %s, accuracy %s, IPC gain %+.0f%%, SPEC §5); paper L1 hit: 11-app average "
                 "(Isolated-Snake %s)" % (pavg[0] if pavg else "n/a", pavg[1] if pavg else "n/a",
                                          pavg[2] if pavg else float("nan"), fig25[2])])

    # cross-checks: final_tier1.md §3, D1_ceilings.md §5, final_tier1.csv
    W = "final_tier1.md §3"
    for _, trows in R.md_tables(os.path.join(RES, "final_tier1.md")):
        if not trows or not trows[0][0].startswith("app / input") or \
                not trows[0][1].startswith("coverage: paper / pred"):
            continue
        for cells in trows[1:]:
            avg = cells[0].startswith("**Average**")
            k = None if avg else t1_key(cells[0], WL)
            if not avg and k is None:
                continue
            r = rows[-1] if avg else rows[WL.index(k)]
            cov, acc, ipc, hit = (R.nums(c) for c in cells[1:5])
            pd = 1 if avg else 0
            for j, doc, dec in ((1, cov[0], pd), (2, cov[1], 1), (3, cov[2], 1), (4, acc[0], pd), (5, acc[1], 1),
                                (6, ipc[0], pd), (7, ipc[1], 1), (9, hit[0], 1), (10, hit[1], 1)):
                r[j].check(doc, dec, W)
            if avg:
                r[11].check(hit[3], 0, W)
                r[12].check(hit[4], 0, W)
    W = "D1_ceilings.md §5 PAPER_V100"
    for head, trows in R.md_tables(os.path.join(RES, "D1_ceilings.md")):
        if not head.startswith("PAPER_V100") or not trows or trows[0][:1] != ["workload"]:
            continue
        for cells in trows[1:]:
            gmr = cells[0].startswith("**GMEAN**")
            k = None if gmr else t1_key(cells[0], WL)
            if not gmr and k is None:
                continue
            r = rows[-1] if gmr else rows[WL.index(k)]
            r[8].check(R.num(cells[3]), 1, W)
            r[7].check(R.num(cells[4]), 1, W)
            r[6].check(R.num(cells[5]), 1 if gmr else 0, W)
    W = "final_tier1.csv"
    for row in csv.DictReader(open(os.path.join(RES, "final_tier1.csv"))):
        k = (row["config"], row["app"], row["args"])
        if row["config"] not in (R.B, R.S) or (k[1], k[2]) not in WL:
            continue
        j = jobs.get(k)
        r = rows[WL.index((k[1], k[2]))]
        check(r[7], j and j["ipc"], float(row["ipc"]), 4, W + " ipc (%s)" % row["config"])
        check(r[9 if row["config"] == R.B else 10], j and j["l1_hit_rate"], float(row["l1_hit_rate"]), 9,
              W + " l1_hit_rate")

    gpu, cap = opts(jobs[(R.S, WL[0][0], WL[0][1])]["path"])
    d["setup"]["s1"] = dict(
        configs="baseline %s; Snake %s; perfect-L1 ceiling %s" % (R.B, R.S, PL1),
        gpu="PAPER_V100 (QV100 config, the Snake paper's V100): " + gpu, cap=cap,
        short="PAPER_V100 (%s)" % gpu,
        traces=trace_arch([j["path"] for j in jobs.values() if j]),
        n="7 apps, %d workloads" % len(WL))
    cols = [R.Col("workload"), R.Col("paper coverage (Fig. 16)", "pct"),
            R.Col("our predicted coverage (coverage_pred)", "pct"),
            R.Col("our delivered coverage (coverage_paper)", "pct"),
            R.Col("paper accuracy, timely (Fig. 17)", "pct"),
            R.Col("our timely accuracy (accuracy_paper)", "pct"), R.Col("paper IPC gain (Fig. 18)", "gain"),
            R.Col("our IPC gain", "gain"), R.Col("perfect-L1 ceiling", "gain"), R.Col("L1 hit baseline", "pct"),
            R.Col("L1 hit Snake", "pct"), R.Col("paper L1 hit baseline (Fig. 25 avg)", "pct"),
            R.Col("paper L1 hit Snake (Fig. 25 avg)", "pct"), R.Col("Status"), R.Col("Notes")]
    desc = ["Slide 1. Our Snake vs the Snake paper (Mostofi et al., MICRO 2023). Tier 1: %s, %s, %s; %s."
            % (d["setup"]["s1"]["gpu"], d["setup"]["s1"]["cap"], "Snake = exact footprint, default config",
               d["setup"]["s1"]["n"]),
            "Coverage and accuracy in % of demand sectors (paper definitions, SPEC OQ-14); IPC gain vs no "
            "prefetching. Definitions: Footnotes; every value's source: Provenance."]
    d["sheets"]["s1"] = (desc, cols, rows)
    d["chart"]["s1"] = dict(labels=[r[0] for r in rows[:-1]] + ["GMEAN\n(7 apps)"],
                            paper=[r[6].v for r in rows], ours=[r[7].v for r in rows], ceil=[r[8].v for r in rows])


# ------------------------------------------------------------ slides 2-4
def sk_paths():
    """{(app, args, label): sim.out}: the SK match runs (sk_match_report
    labels) and the base QV100 no-prefetch runs of the SK_SC stage ("base")."""
    want = {c: lab for lab, src, c in SKM_CONFIGS if src == "skm"}
    out = {}
    for p in glob.glob(os.path.join(ROOT, SKC, "*", "runs", "*", "*", Q, "sim.out")):
        app, args = p.split(os.sep)[-4:-2]
        out[(app, args, "base")] = p
    for p in glob.glob(os.path.join(ROOT, SKM, "runs", "*", "*", "*", "sim.out")):
        app, args, cfg = p.split(os.sep)[-4:-1]
        if cfg in want:
            out[(app, args, want[cfg])] = p
    return out


def sc_values():
    """SnakeCharmer's SK averages from results/SK_SC.md §4."""
    txt = R.rd(os.path.join(RES, "SK_SC.md"))
    vals = {}
    for t in R.md_tables_text(txt):
        if t and t[0][0] == "metric" and t[0][1].startswith("SnakeCharmer SK"):
            key = {"coverage": "cov", "accuracy": "acc", "on time": "on", "too early": "early", "too late": "late"}
            vals = {key[c[0]]: float(c[1]) for c in t[1:] if c[0] in key}
    m = re.search(r"the AVG over its (\d+) apps", txt)
    vals["n"] = int(m.group(1)) if m else None
    vals["no_ipc"] = "SnakeCharmer reports none for SK" in txt
    for k, v in BRIEF_SC.items():
        if k in vals:
            check(None, vals[k], v, 1, "SK_SC.md §4 vs the task brief (%s)" % k)
    return vals


def slide2_4(d):
    paths = sk_paths()
    runs = {}
    for k, p in paths.items():
        s = parse(p)
        ok = s["_done"] and (k[2] == "base" or (
            s.get("_pf_invariant") == "OK" and s.get("_funnel_invariant") == "OK"
            and s.get("_funnel_src_invariant") == "OK" and s.get("pf_woken_warps", 0) == 0))
        CHECKS.append(("run finished, invariants OK: " + rel(p), ok))
        if not ok:
            MISMATCH.append(("run check", rel(p), "unfinished or invariant broken"))
        runs[k] = (p, s, set(build(p).split()))
    wl = sorted({(a, g) for a, g, lab in runs if lab == "SK"}, key=lambda w: (w[0] not in SEVEN, w[0], w[1]))
    per = {}  # (app, args, label) -> metrics + gain + job info
    for a, g in wl:
        bp, bs, bt = runs[(a, g, "base")]
        for lab, _ in SK_LABELS:
            if (a, g, lab) not in runs:
                continue
            p, s, t = runs[(a, g, lab)]
            for pp, tt in ((p, t), (bp, bt)):
                use(pp, tt)
            m = {k: nan2none(v) for k, v in metrics(s).items()}
            m.update(gain=(s["_ipc"] / bs["_ipc"] - 1) * 100, ipc=s["_ipc"], ipc_base=bs["_ipc"],
                     src="%s; %s" % (rel(os.path.dirname(p)), rel(os.path.dirname(bp))), tags=t | bt)
            per[(a, g, lab)] = m
    sc5 = [w for w in wl if w[0] not in SEVEN]
    seven = [w for w in wl if w[0] in SEVEN]
    scopes = [("SnakeCharmer's %d workloads" % len(sc5), sc5, "SnakeCharmer apps"),
              ("all %d workloads" % len(wl), wl, "all")]
    cfg_of = {lab: c for lab, src, c in SKM_CONFIGS}
    skc_sets = sorted({run_dir_of(runs[(a, g, "base")][0]) for a, g in wl})
    srcs = lambda lab, ws: "%s/runs/*/*/%s (%d jobs); no prefetching: %s/runs/*/*/%s" % (
        SKM, cfg_of[lab], len(ws), " + ".join(skc_sets), Q)
    HOW = {"cov_on": "sc_ld_ontime / (sc_ld_ontime + sc_ld_miss)",
           "cov_late": "(sc_ld_ontime + sc_ld_late) / (sc_ld_ontime + sc_ld_miss)",
           "acc": "(sc_ld_ontime + sc_pf_fills - pf_filled) / sc_pf_fills",
           "t_on": "sc_ld_ontime / (sc_ld_ontime + sc_ld_late + sc_ld_early)",
           "t_early": "sc_ld_early / (sc_ld_ontime + sc_ld_late + sc_ld_early)",
           "t_late": "sc_ld_late / (sc_ld_ontime + sc_ld_late + sc_ld_early)",
           "pfev": "sc_victim_demand_by_pf / sc_victim_sectors"}

    def mean_v(lab, ws, key):
        xs = [per[w + (lab,)][key] for w in ws if (w + (lab,)) in per]
        vals = [x for x in xs if x is not None]
        return V(sum(vals) / len(vals) if vals else None, srcs(lab, ws),
                 "%s per workload; arithmetic mean over the %d of %d workloads with a value"
                 % (HOW[key], len(vals), len(ws)),
                 set().union(*[per[w + (lab,)]["tags"] for w in ws if (w + (lab,)) in per])), len(vals)

    def gmean_v(lab, ws, what):
        xs = [per[w + (lab,)]["gain"] for w in ws if (w + (lab,)) in per]
        return V(gm(xs), srcs(lab, ws), "GMEAN over %d workloads%s of gpu_tot_ipc %s / no prefetching - 1"
                 % (len(xs), what, lab), set().union(*[per[w + (lab,)]["tags"] for w in ws if (w + (lab,)) in per]))

    # ---- slide 2
    sc = sc_values()
    SCS = "docs/snake/results/SK_SC.md §4 (SnakeCharmer Fig. 12 AVG and Fig. 4, A100)"
    rows = [["SnakeCharmer SK (A100)", "%s apps, A100 average" % sc.get("n"),
             V(sc.get("n"), SCS, "number of apps in SnakeCharmer's average"),
             V(sc.get("cov"), SCS, "SnakeCharmer coverage = H / (H + M)"),
             V(None, SCS, "SnakeCharmer reports one coverage"),
             V(sc.get("acc"), SCS, "SnakeCharmer accuracy"), V(sc.get("on"), SCS, "SnakeCharmer timeliness"),
             V(sc.get("early"), SCS, "SnakeCharmer timeliness"), V(sc.get("late"), SCS, "SnakeCharmer timeliness"),
             V(BRIEF_SC["evict"], "task brief (2026-10-27); not in our docs",
               "SnakeCharmer's prefetch-induced eviction share as given in the brief"),
             V("not reported", "docs/snake/results/SK_SC.md §5 heading" if sc.get("no_ipc") else "task brief",
               "SnakeCharmer reports no IPC for SK"),
             PUB, "coverage incl. late: n/a, SnakeCharmer reports one coverage (H / (H + M), late uses not "
             "separated; SPEC OQ-36). Eviction share: from the task brief, its definition not checked against "
             "ours. Full runs, zero-compressed cache"]]
    means = {}
    for scope, ws, _ in scopes:
        for lab, name in SK_LABELS:
            vs = {}
            ns = {}
            for key in ("cov_on", "cov_late", "acc", "t_on", "t_early", "t_late", "pfev"):
                vs[key], ns[key] = mean_v(lab, ws, key)
            vs["ipc"] = gmean_v(lab, ws, "")
            means[(lab, scope)] = vs
            short_n = [k for k in ("acc", "t_on", "pfev") if ns[k] < len(ws)]
            note = "means over workloads (IPC: GMEAN)"
            if short_n:
                note += "; n with a value: accuracy %d, timeliness %d, evictions %d of %d (no prefetch used or " \
                        "no eviction: '-' in SK_match_qv100.md)" % (ns["acc"], ns["t_on"], ns["pfev"], len(ws))
            rows.append([name, scope, V(len(ws), srcs(lab, ws), "workloads in the scope"), vs["cov_on"],
                         vs["cov_late"], vs["acc"], vs["t_on"], vs["t_early"], vs["t_late"], vs["pfev"], vs["ipc"],
                         PRE, note])
    W = "SK_match_qv100.md Means"
    smap = {s[2]: s[0] for s in scopes}
    for t in R.md_tables_text(R.rd(os.path.join(RES, "SK_match_qv100.md"))):
        if not t or t[0][:3] != ["set", "config", "n"]:
            continue
        for c in t[1:]:
            if c[0] == "SnakeCharmer":
                for j, key in ((4, "cov"), (6, "acc"), (7, "on"), (8, "early"), (9, "late")):
                    check(rows[0][{"cov": 3, "acc": 5, "on": 6, "early": 7, "late": 8}[key]],
                          sc.get(key), R.num(c[j]), 1, W + " SnakeCharmer row")
            if c[0] not in smap or (c[1], smap[c[0]]) not in means:
                continue
            vs = means[(c[1], smap[c[0]])]
            check(None, len([w for w in (sc5 if c[0] != "all" else wl)]), R.num(c[2]), 0, W + " n")
            vs["ipc"].check(R.num(c[3]), 2, W)
            for j, key in ((4, "cov_on"), (5, "cov_late"), (6, "acc"), (7, "t_on"), (8, "t_early"),
                           (9, "t_late"), (10, "pfev")):
                vs[key].check(R.num(c[j]), 1, W)
    cols = [R.Col("config"), R.Col("scope"), R.Col("workloads", "int"),
            R.Col("coverage_sc, on-time only", "pct"), R.Col("coverage_sc, incl. late", "pct"),
            R.Col("accuracy_sc", "pct"), R.Col("on time", "pct"), R.Col("too early", "pct"),
            R.Col("too late", "pct"), R.Col("prefetch-induced eviction share", "pct"),
            R.Col("IPC vs no prefetching (GMEAN)", "gain"), R.Col("Status"), R.Col("Notes")]
    gpu, cap = opts(os.path.join(ROOT, per[sc5[0] + (BEST,)]["src"].split("; ")[0]))
    d["setup"]["s2"] = dict(
        configs="; ".join("%s = %s" % (lab, cfg_of[lab]) for lab, _ in SK_LABELS) + "; no prefetching = " + Q,
        gpu="base QV100 (SnakeCharmer's V100 setup): " + gpu, cap=cap,
        short="base QV100 (%s)" % gpu,
        traces=trace_arch([os.path.join(ROOT, per[w + (BEST,)]["src"].split("; ")[0]) for w in wl]),
        n="%d workloads: our 7 apps (%d workloads) and SnakeCharmer's apps without bfs (%d)"
          % (len(wl), len(seven), len(sc5)))
    desc = ["Slide 2. Snake as SnakeCharmer's SK (Tang et al., MICRO 2026) vs SnakeCharmer's published SK "
            "averages. Ours: %s, %s." % (d["setup"]["s2"]["gpu"], d["setup"]["s2"]["cap"]),
            "SnakeCharmer's metrics per 32 B sector (SPEC OQ-36); on time / too early / too late are shares of "
            "used prefetches. Means are arithmetic over workloads, IPC a GMEAN."]
    d["sheets"]["s2"] = (desc, cols, rows)
    sc_row = rows[0]
    best5 = means[(BEST, scopes[0][0])]
    best14 = means[(BEST, scopes[1][0])]
    d["chart"]["s2"] = dict(bars=[("SnakeCharmer SK\n(A100, %s apps)" % sc.get("n"),
                                   [sc_row[6].v, sc_row[7].v, sc_row[8].v]),
                                  ("%s\n(SnakeCharmer's %d)" % (BEST, len(sc5)),
                                   [best5["t_on"].v, best5["t_early"].v, best5["t_late"].v]),
                                  ("%s\n(all %d)" % (BEST, len(wl)),
                                   [best14["t_on"].v, best14["t_early"].v, best14["t_late"].v])])

    # ---- slide 3 (derived)
    rows3 = []
    for name, scope, acc, on, late, early, srcname in (
            ("SnakeCharmer SK (A100)", sc_row[1], sc_row[5], sc_row[6], sc_row[8], sc_row[7], "SnakeCharmer SK row"),
            (dict(SK_LABELS)[BEST], scopes[0][0], best5["acc"], best5["t_on"], best5["t_late"], best5["t_early"],
             "%s, %s row" % (BEST, scopes[0][0]))):
        S2 = "sheet '%s', %s" % (NAMES["s2"], srcname)
        dv = lambda x, y, what: V(x.v * y.v / 100 if x.v is not None and y.v is not None else None, S2,
                                  "accuracy_sc x %s share / 100 = %.1f x %.1f / 100" % (what, x.v, y.v),
                                  x.tags | y.tags)
        rows3.append([name, scope, V(acc.v, S2, "used = accuracy_sc", acc.tags), dv(acc, on, "on-time"),
                      dv(acc, late, "too-late"), dv(acc, early, "too-early"),
                      V(100 - acc.v if acc.v is not None else None, S2, "wasted = 100 - accuracy_sc", acc.tags),
                      DER, "derived, approximate; definitions differ (Footnotes): the shares are of used "
                      "prefetches, our too-early sectors are not counted in our accuracy_sc, and products of means "
                      "are not means of products"])
    cols3 = [R.Col("config"), R.Col("scope"), R.Col("used", "dec1"), R.Col("used on time", "dec1"),
             R.Col("used late", "dec1"), R.Col("used early", "dec1"), R.Col("wasted", "dec1"),
             R.Col("Status"), R.Col("Notes")]
    desc3 = ["Slide 3. Per 100 prefetched sectors (fills): derived from sheet '%s'. Derived, approximate; "
             "definitions differ." % NAMES["s2"],
             "used = accuracy_sc; used on time / late / early = accuracy_sc x on-time / too-late / too-early "
             "share; wasted = 100 - used."]
    d["sheets"]["s3"] = (desc3, cols3, rows3)
    d["chart"]["s3"] = dict(bars=[(r[0].split(" (")[0] + ("\n(A100, %s apps)" % sc.get("n") if i == 0 else
                                                          "\n(SnakeCharmer's %d)" % len(sc5)),
                                   [r[3].v, r[4].v, r[5].v, r[6].v]) for i, r in enumerate(rows3)])

    # ---- slide 4
    rows4 = []
    W = "SK_match_qv100.md Per workload"
    doc_ipc = {}
    for t in R.md_tables_text(R.rd(os.path.join(RES, "SK_match_qv100.md"))):
        if t and t[0][:3] == ["workload", "config", "IPC vs no prefetch"]:
            for c in t[1:]:
                doc_ipc[(c[0], c[1])] = R.num(c[2]) if c[2] else None
    for grp, ws in (("our 7 apps", seven), ("SnakeCharmer apps", sc5)):
        for a, g in sorted(ws, key=lambda w: -per[w + (BEST,)]["gain"]):
            m = per[(a, g, BEST)]
            name = short(a, g) + (" †" if a.startswith(INPUT_DIFFERS) else "")
            notes = []
            if a.startswith(INPUT_DIFFERS):
                notes.append("† input differs from SnakeCharmer Table II (SK_SC.md §1)")
            if a in MISDECODED:
                notes.append("%s of the trace's global memory lines mis-decoded (SK_SC.md §1)" % MISDECODED[a])
            if a == NAIVE:
                notes.append("dominates the means: latency-bound, low occupancy (SK_SC.md §5)")
            jb, jbase = m["src"].split("; ")
            row = [name, grp, V(m["ipc_base"], jbase, "gpu_tot_ipc, no prefetching", m["tags"]),
                   V(m["ipc"], jb, "gpu_tot_ipc, %s" % BEST, m["tags"]),
                   V(m["gain"], m["src"], "gpu_tot_ipc %s / no prefetching - 1" % BEST, m["tags"]),
                   PRE, "; ".join(notes)]
            row[4].check(doc_ipc.get((short(a, g) + (" †" if a.startswith(INPUT_DIFFERS) else ""), BEST)), 2, W)
            rows4.append(row)
    nonaive = [w for w in wl if w[0] != NAIVE]
    gm_rows = [("GMEAN, all %d workloads" % len(wl), wl, "all", ""),
               ("GMEAN without particlefilter_naive (%d)" % len(nonaive), nonaive, None,
                "not in the results docs; computed here from the same runs"),
               ("GMEAN, our 7 apps (%d workloads)" % len(seven), seven, "7 apps", ""),
               ("GMEAN, SnakeCharmer apps (%d workloads)" % len(sc5), sc5, "SnakeCharmer apps", "")]
    doc_gm = {}
    for t in R.md_tables_text(R.rd(os.path.join(RES, "SK_match_qv100.md"))):
        if t and t[0][:3] == ["set", "config", "n"]:
            for c in t[1:]:
                if c[1] == BEST:
                    doc_gm[c[0]] = R.num(c[3])
    for name, ws, docset, note in gm_rows:
        v = gmean_v(BEST, ws, "")
        if docset:
            v.check(doc_gm.get(docset), 2, "SK_match_qv100.md Means (%s)" % docset)
        rows4.append([name, "summary", "", "", v, PRE, note or "GMEAN over workloads"])
    rows4.append(["bfs (excluded)", "SnakeCharmer apps", None, None, None, None,
                  "excluded: %s of its global memory lines are mis-decoded after the parser fix, a tracer bug the "
                  "fix cannot repair (SK_SC.md §1); no SK match runs" % MISDECODED.get("bfs-rodinia-3.1", "n/a")])
    cols4 = [R.Col("workload"), R.Col("group"), R.Col("IPC, no prefetching", "ipc"), R.Col("IPC, %s" % BEST, "ipc"),
             R.Col("IPC gain", "gain"), R.Col("Status"), R.Col("Notes")]
    desc4 = ["Slide 4. Best Snake variant (%s: SK settings, same-CTA inter-warp targets, per-warp Head, miss-queue "
             "reserve 0, decoupling off) vs no prefetching; %s, %s." % (BEST, d["setup"]["s2"]["gpu"],
                                                                       d["setup"]["s2"]["cap"]),
             "Workloads sorted by IPC gain within each group. † input differs from SnakeCharmer Table II. bfs "
             "excluded (mis-decoded trace lines)."]
    d["sheets"]["s4"] = (desc4, cols4, rows4)
    d["chart"]["s4"] = dict(bars=[(r[0], r[1], r[4].v) for r in rows4 if r[1] in ("our 7 apps", "SnakeCharmer apps")
                                  and isinstance(r[4], V) and r[4].v is not None],
                            gmeans=[(r[0], r[4].v) for r in rows4 if r[1] == "summary"])
    d["gm_naive"] = (rows4[-5][4].v, rows4[-4][4].v)


# ------------------------------------------------------------ footnotes and provenance
def footnotes(d):
    S = "docs/snake/SPEC.md"
    s1, s2 = d["setup"]["s1"], d["setup"]["s2"]
    g14, g13 = d["gm_naive"]
    rows = [
        ["metric", "IPC gain", "gpu_tot_ipc(config) / gpu_tot_ipc(no prefetching) - 1, same workload and setup",
         "sim.out of both runs"],
        ["metric", "GMEAN", "exp(mean(ln(1 + gain))) - 1. Slide 1: over the 7 apps, an app's inputs (hotspot, lud) "
         "combined first; slides 2 and 4: over workloads", "final_tier1.md §3; SK_match_qv100.md Means"],
        ["metric", "paper coverage (Fig. 16)", "'correctly predicted addresses / total demand addresses' (Snake "
         "paper §4); values read off the bar chart, ±3 pp", S + " §5"],
        ["metric", "our predicted coverage (coverage_pred)", "demand sectors that an earlier prediction named, "
         "issued or not, within snake_pred_window cycles / demand sectors (snake_demand_pred_matched / "
         "snake_demand_sector_accesses)", S + " OQ-14"],
        ["metric", "our delivered coverage (coverage_paper)", "(timely + late prefetched sectors) / demand sectors "
         "((snake_pf_useful_timely + snake_pf_late) / snake_demand_sector_accesses)", S + " OQ-14"],
        ["metric", "paper accuracy, timely (Fig. 17)", "'timely correctly predicted addresses / total demand "
         "addresses': timely coverage, not useful / issued; read off the bar chart, ±3 pp", S + " §5, OQ-14"],
        ["metric", "our timely accuracy (accuracy_paper)", "timely prefetched sectors / demand sectors "
         "(snake_pf_useful_timely / snake_demand_sector_accesses)", S + " OQ-14"],
        ["metric", "perfect-L1 ceiling", "IPC gain with -gpgpu_perfect_l1d 1: every L1-cached global load hits at L1 "
         "latency. An upper bound: it also removes the loads' memory traffic, which a prefetcher still moves",
         "results/D1_ceilings.md §3"],
        ["metric", "L1 hit rate", "demand L1D hit rate: Total_core_cache_stats_breakdown[GLOBAL_ACC_R][HIT] / "
         "[TOTAL_ACCESS]; the paper's is the Fig. 25 average over its 11 apps", "final_tier1.md §3; " + S + " §5"],
        ["metric", "coverage_sc, on-time only", "on-time uses / (on-time + miss): sc_ld_ontime / "
         "(sc_ld_ontime + sc_ld_miss); SnakeCharmer's H / (H + M). sc_ld_miss already holds the late uses "
         "(every MISS return, MSHR merges included; debug plan S5)",
         S + " OQ-36"],
        ["metric", "coverage_sc, incl. late", "(on-time + late) / (on-time + miss): the same with late uses "
         "in H. SnakeCharmer does not say which it counts", S + " OQ-36"],
        ["metric", "accuracy_sc", "(on-time uses + fills a load had merged into) / prefetch fills: (sc_ld_ontime + "
         "sc_pf_fills - pf_filled) / sc_pf_fills", S + " OQ-36"],
        ["metric", "on time / too early / too late", "shares of used prefetches (first load to a prefetched "
         "sector): hit after the fill / evicted unused earlier in the kernel / merged into the in-flight MSHR entry",
         S + " OQ-36"],
        ["metric", "prefetch-induced eviction share", "demand-filled victim sectors evicted by a prefetch allocation "
         "/ all victim sectors (sc_victim_demand_by_pf / sc_victim_sectors). SnakeCharmer's %.1f is from the task "
         "brief; its definition was not checked against ours" % BRIEF_SC["evict"], S + " OQ-36; task brief"],
        ["metric", "per 100 prefetched sectors", "used = accuracy_sc; used on time / late / early = accuracy_sc x "
         "on-time / too-late / too-early share / 100; wasted = 100 - used. Derived from the means, approximate",
         "sheet '%s'" % NAMES["s3"]],
        ["status", VER, "1B-instruction Tier 1 run (final evaluation), checked against the results docs",
         "final_tier1.md"],
        ["status", PRE, "100M-instruction run (SK match Step 1), checked against the results docs",
         "SK_match_qv100.md"],
        ["status", DER, "computed here from values in other sheets", "sheet '%s'" % NAMES["s3"]],
        ["status", PUB, "a published average (SnakeCharmer), not reproduced by us", "SK_SC.md §4"],
        ["setup", "slide 1: configs", s1["configs"], "s5_runs.sh tier1 list; D1_ceilings.md §4"],
        ["setup", "slide 1: GPU model", s1["gpu"], "gpgpusim.config of the runs"],
        ["setup", "slide 1: traces", s1["traces"] + " (accel-sim traces of V100 runs, '-binary version' in kernel-1)",
         "hw_run/*/traces"],
        ["setup", "slide 1: instruction cap", s1["cap"], "1B_INSN token"],
        ["setup", "slide 1: workloads", s1["n"], "final_tier1.md"],
        ["setup", "slides 2-4: configs", s2["configs"], "SK_match_qv100.md Step 1"],
        ["setup", "slides 2-4: GPU model", s2["gpu"], "gpgpusim.config of the runs; SK_SC.md §1"],
        ["setup", "slides 2-4: traces", s2["traces"], "hw_run/*/traces"],
        ["setup", "slides 2-4: instruction cap", s2["cap"], "100M_INSN token"],
        ["setup", "slides 2-4: workloads", s2["n"], "SK_match_qv100.md"],
        ["caveat", "SnakeCharmer's numbers", "A100 averages over its %s apps, full runs, with a zero-compressed "
         "cache. Ours: the V100 model at 100M instructions, %s" % (d["sc_n"], s2["n"]), "SK_SC.md §4, §8"],
        ["caveat", "accuracy definitions differ", "the Snake paper's 'accuracy' is timely coverage (per demand "
         "sector); SnakeCharmer's is per prefetch fill (accuracy_sc); our too-early sectors are outside accuracy_sc",
         S + " OQ-14, OQ-36"],
        ["caveat", "particlefilter_naive dominates the means", "%s IPC GMEAN over all workloads %+.2f%%, without "
         "particlefilter_naive %+.2f%%; it also dominates arithmetic means of the rates" % (BEST, g14, g13),
         "sheet '%s'" % NAMES["s4"]],
        ["caveat", "no real A100 traces", "no sm_80 traces exist locally; every run uses V100 (sm_70) traces, so "
         "SnakeCharmer's A100 setup was not run", "SK_SC.md §2"],
        ["caveat", "paper values", "Snake paper values are read off bar charts (±3 pp); its averages are recomputed "
         "over our 7 apps", S + " §5"],
        ["caveat", "inputs and traces", "heartwall and streamcluster inputs differ from SnakeCharmer Table II; bfs "
         "is excluded (%s of global memory lines mis-decoded); streamcluster has %s" % (
             MISDECODED.get("bfs-rodinia-3.1"), MISDECODED.get("streamcluster-rodinia-2.0-ft")), "SK_SC.md §1"],
    ]
    cols = [R.Col("section"), R.Col("item"), R.Col("definition / formula / text"), R.Col("source")]
    d["sheets"]["foot"] = (["Footnotes: metric definitions with formula and source, the setup per slide, status "
                            "values and caveats."], cols, rows)


def provenance(d, names, pngs):
    rows = []
    for key in ("s1", "s2", "s3", "s4"):
        desc, cols, trows = d["sheets"][key]
        for r in trows:
            label = r[0] + ("" if key == "s1" else " | " + str(r[1]))
            for c, x in zip(cols, r):
                if isinstance(x, V):
                    rows.append(["value", names[key], label, c.name, x.v, x.src, x.how,
                                 " ".join(sorted(t for t in x.tags if t)) or "", "; ".join(x.chk) or "",
                                 r[-2], ""])
    for dr, tags in sorted(USED.items()):
        n = sum(len(j) for j in tags.values())
        rows.append(["run dir", dr, "", "", n, "%d jobs used" % n, "",
                     "; ".join("%s (%d jobs)" % (k, len(j)) for k, j in sorted(tags.items())), "", "",
                     "all build tags _modified_0.0" if all(k.count("_modified_0.0") == 2 for k in tags)
                     else "CHECK: a build tag is not _modified_0.0"])
    for f in ("docs/snake/results/final_tier1.md", "docs/snake/results/final_tier1.csv",
              "docs/snake/results/D1_ceilings.md", "docs/snake/results/SK_SC.md",
              "docs/snake/results/SK_match_qv100.md", "docs/snake/SPEC.md"):
        rows.append(["source file", f, "", "", "", "", "last commit", git("log", "-1", "--format=%h %ad %s",
                                                                          "--date=short", "--", f), "", "", ""])
    for repo, path in (("accel-sim", ROOT), ("gpgpu-sim", os.path.join(ROOT, "gpu-simulator", "gpgpu-sim"))):
        rows.append(["tree", repo, "", "", "", "", "HEAD when generated",
                     subprocess.run(["git", "-C", path, "log", "-1", "--format=%h (%D)"], capture_output=True,
                                    text=True).stdout.strip(), "", "", ""])
    for name in sorted(pngs):
        rows.append(["chart", name, "", "", "", "values of sheet '%s'" % names[CHART_SHEET[name]], "", "", "", "",
                     ""])
    by = {}
    for where, ok in CHECKS:
        w = where.split(": ")[0] if where.startswith("run finished") else where
        by.setdefault(w, [0, 0])
        by[w][0] += 1
        by[w][1] += 0 if ok else 1
    for w, (n, bad) in sorted(by.items()):
        rows.append(["cross-check", w, "", "", n, "%d compared, %d mismatches" % (n, bad), "", "", "", "", ""])
    for where, ours, doc in MISMATCH:
        rows.append(["MISMATCH", where, "", "", "", "ours %s vs doc %s" % (ours, doc), "", "", "", "", ""])
    cols = [R.Col("kind"), R.Col("sheet / item"), R.Col("row"), R.Col("column"), R.Col("value", "dec2"),
            R.Col("source (file / run dirs)"), R.Col("computed as"), R.Col("build tags / commit"),
            R.Col("check vs results docs"), R.Col("Status"), R.Col("Notes")]
    return (["Provenance: every value of the slide sheets with its source, formula, build tags and check against "
             "the results docs; then the run dirs, source-file commits, trees and cross-check totals.",
             "Generated by util/snake/make_slides_data.py from the accel-sim tree at %s."
             % git("rev-parse", "--short", "HEAD")], cols, rows)


def collect():
    del CHECKS[:], MISMATCH[:]
    USED.clear()
    d = {"sheets": {}, "chart": {}, "setup": {}}
    slide1(d)
    slide2_4(d)
    d["sc_n"] = d["sheets"]["s2"][2][0][2].v
    footnotes(d)
    return d


# ------------------------------------------------------------ charts
# Reference palette of the dataviz skill (slots 1-4, its validated order),
# light surface; a neutral for "wasted".
INK, INK2, GRID, AXIS = "#0b0b0b", "#52514e", "#e1e0d9", "#c3c2b7"
SERIES = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100"]
NEUTRAL = "#898781"
CHART_SHEET = {}


def style():
    plt.rcParams.update({"font.size": 16, "axes.labelsize": 17, "xtick.labelsize": 15, "ytick.labelsize": 16,
                         "legend.fontsize": 16, "axes.edgecolor": AXIS, "axes.labelcolor": INK2,
                         "xtick.color": INK2, "ytick.color": INK, "text.color": INK,
                         "font.family": "DejaVu Sans"})


def frame(title, subtitle, top=0.80):
    fig, ax = plt.subplots(figsize=(13.33, 7.5))
    fig.subplots_adjust(left=0.08, right=0.97, top=top, bottom=0.14)
    fig.text(0.04, 0.95, title, fontsize=25, weight="bold", va="top")
    fig.text(0.04, 0.885, subtitle, fontsize=14, color=INK2, va="top", wrap=True)
    for sp in ("top", "right"):
        ax.spines[sp].set_visible(False)
    ax.set_axisbelow(True)
    return fig, ax


def png(fig):
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150, facecolor="white")
    plt.close(fig)
    return buf.getvalue()


def stacked(ax, bars, names, colors, unit):
    ys = list(range(len(bars)))[::-1]
    for (label, vals), y in zip(bars, ys):
        left = 0
        for v, c in zip(vals, colors):
            v = v or 0
            ax.barh(y, v, left=left, color=c, edgecolor="white", linewidth=2, height=0.62)
            if v >= 4.5:
                ax.text(left + v / 2, y, "%.1f%s" % (v, unit), ha="center", va="center", fontsize=17,
                        weight="bold", color=INK if c == SERIES[3] else "white")
            left += v
    ax.set_yticks(ys)
    ax.set_yticklabels([b[0] for b in bars])
    ax.set_xlim(0, 100)
    ax.xaxis.grid(True, color=GRID, lw=0.8)
    ax.spines["left"].set_visible(False)
    ax.tick_params(axis="y", length=0)
    handles = [plt.Rectangle((0, 0), 1, 1, color=c) for c in colors]
    ax.legend(handles, names, loc="lower center", bbox_to_anchor=(0.5, 1.0), ncol=len(names), frameon=False)


def charts(d, date=""):
    style()
    out = {}
    sfx = ("_" + date) if date else ""
    s1, su1 = d["chart"]["s1"], d["setup"]["s1"]
    fig, ax = frame("IPC gain: Snake paper vs our Snake vs perfect L1",
                    "Tier 1, %s, %s, Snake exact footprint. Paper values read off Fig. 18 (±3 pp). "
                    "GMEAN over the 7 apps." % (su1["short"], su1["cap"].split(" (")[0]))
    n, w = len(s1["labels"]), 0.27
    series = (("Snake paper (Fig. 18)", s1["paper"]), ("our Snake", s1["ours"]), ("perfect-L1 ceiling", s1["ceil"]))
    for i, (name, vals) in enumerate(series):
        xs = [k + (i - 1) * w for k in range(n)]
        ax.bar(xs, [v or 0 for v in vals], w * 0.92, color=SERIES[i], label=name)
        for k, (x, v) in enumerate(zip(xs, vals)):
            if v is None:
                ax.text(x, 0.5, "n/a", ha="center", fontsize=12, color=INK2, rotation=90)
            elif i == 1 or k == n - 1:
                ax.text(x, v + (0.8 if v >= 0 else -0.8), "%+.1f" % v, ha="center", va="bottom" if v >= 0 else "top",
                        fontsize=13, color=INK, weight="bold" if i == 1 else "normal")
    ax.axhline(0, color=AXIS, lw=1.2)
    ax.axvline(n - 1.5, color=AXIS, lw=1, ls=":")
    ax.set_xticks(range(n))
    ax.set_xticklabels([l if "\n" in l else l.replace(" ", "\n", 1) for l in s1["labels"]])
    ax.set_ylabel("IPC gain vs no prefetching (%)")
    ax.yaxis.grid(True, color=GRID, lw=0.8)
    ax.legend(loc="upper right", frameon=False, ncol=3)
    out["slide1_ipc_vs_paper%s.png" % sfx] = png(fig)
    CHART_SHEET["slide1_ipc_vs_paper%s.png" % sfx] = "s1"

    su2 = d["setup"]["s2"]
    early0 = [b[0].replace("\n", " ") for b in d["chart"]["s2"]["bars"][1:] if (b[1][1] or 0) == 0]
    fig, ax = frame("Prefetch timeliness: SnakeCharmer's SK vs our best Snake",
                    "Share of used prefetches. SnakeCharmer: published A100 average. Ours: %s, %s; mean over "
                    "workloads.%s" % (su2["short"], su2["cap"].split(" (")[0],
                                      (" Too early = 0.0%% for %s." % " and ".join(early0)) if early0 else ""))
    stacked(ax, d["chart"]["s2"]["bars"], ["on time", "too early", "too late"], SERIES[:3], "%")
    ax.set_xlabel("% of used prefetches")
    fig.subplots_adjust(left=0.24, top=0.74)
    out["slide2_timeliness%s.png" % sfx] = png(fig)
    CHART_SHEET["slide2_timeliness%s.png" % sfx] = "s2"

    fig, ax = frame("Per 100 prefetched sectors",
                    "Derived, approximate; definitions differ: used = accuracy_sc, split by the timeliness shares; "
                    "wasted = 100 - used.")
    stacked(ax, d["chart"]["s3"]["bars"], ["used on time", "used late", "used early", "wasted"],
            [SERIES[0], SERIES[2], SERIES[1], NEUTRAL], "")
    ax.set_xlabel("prefetched sectors (fills), out of 100")
    fig.subplots_adjust(left=0.24, bottom=0.2, top=0.74)
    out["slide3_per100_fills%s.png" % sfx] = png(fig)
    CHART_SHEET["slide3_per100_fills%s.png" % sfx] = "s3"

    s4 = d["chart"]["s4"]
    bars = sorted(s4["bars"], key=lambda b: b[2])
    gtxt = "; ".join("%s %+.1f%%" % (k.replace("GMEAN, ", "").replace("GMEAN ", ""), v) for k, v in s4["gmeans"])
    fig, ax = frame("%s vs no prefetching: IPC gain per workload" % BEST,
                    "%s, %s. GMEAN: %s. bfs excluded (mis-decoded trace lines). † input differs from "
                    "SnakeCharmer Table II." % (su2["short"], su2["cap"].split(" (")[0], gtxt), top=0.78)
    grp_color = {"our 7 apps": SERIES[0], "SnakeCharmer apps": SERIES[1]}
    xmax = max(b[2] for b in bars)
    for y, (label, grp, v) in enumerate(bars):
        ax.barh(y, v, color=grp_color[grp], height=0.7)
        inside = v > 0.6 * xmax
        ax.text(v - 0.6 if inside else (v + 0.6 if v >= 0 else v - 0.6), y, "%+.1f%%" % v, va="center",
                ha="right" if inside or v < 0 else "left", fontsize=14,
                color="white" if inside else INK, weight="bold" if inside else "normal")
    ax.set_yticks(range(len(bars)))
    ax.set_yticklabels([b[0] for b in bars], fontsize=15)
    ax.axvline(0, color=AXIS, lw=1.2)
    ax.set_xlim(min(-10, min(b[2] for b in bars) - 9), xmax * 1.05)
    ax.xaxis.grid(True, color=GRID, lw=0.8)
    ax.tick_params(axis="y", length=0)
    ax.spines["left"].set_visible(False)
    ax.set_xlabel("IPC gain vs no prefetching (%)")
    ax.legend([plt.Rectangle((0, 0), 1, 1, color=c) for c in grp_color.values()], list(grp_color),
              loc="lower right", frameon=False)
    fig.subplots_adjust(left=0.22, top=0.78, bottom=0.11)
    out["slide4_best_snake_ipc%s.png" % sfx] = png(fig)
    CHART_SHEET["slide4_best_snake_ipc%s.png" % sfx] = "s4"
    return out


# ------------------------------------------------------------ workbook
def add_sheets(wb, d, names, pngs=None):
    """Writes the slide, Footnotes and Provenance sheets (sheet names from
    names); pngs (file name -> PNG bytes) are placed under their tables."""
    pngs = pngs or {}
    for key in ("s1", "s2", "s3", "s4", "foot"):
        desc, cols, rows = d["sheets"][key]
        plain = [[x.v if isinstance(x, V) else x for x in r] for r in rows]
        imgs = [b for f, b in sorted(pngs.items()) if CHART_SHEET.get(f) == key]

        def extra(ws, hr, imgs=imgs, n=len(plain)):
            for b in imgs:
                img = XLImage(io.BytesIO(b))
                img.width, img.height = 1000, 562
                ws.add_image(img, "A%d" % (hr + n + 3))
        R.write_sheet(wb, names[key], desc, cols, plain, extra)
    desc, cols, rows = provenance(d, names, pngs)
    R.write_sheet(wb, names["prov"], desc, cols, rows)


def verify(path, d, names):
    """Reopen the workbook: sheets, row counts, five values vs their sources."""
    ok = True
    wb = load_workbook(path)
    print("## Verification (reopened %s)" % rel(path))
    for key in ("s1", "s2", "s3", "s4", "foot"):
        desc, cols, rows = d["sheets"][key]
        ws = wb[names[key]]
        hr = len(desc) + 1
        got = ws.max_row - hr
        heads = [ws.cell(row=hr, column=j + 1).value for j in range(len(cols))]
        good = got == len(rows) and heads == [c.name for c in cols]
        ok &= good
        print("- sheet '%s': %d data rows (expected %d), header %s" % (names[key], got, len(rows),
                                                                       "ok" if heads == [c.name for c in cols] else "WRONG"))
    prov = wb[names["prov"]]
    print("- sheet '%s': %d rows" % (names["prov"], prov.max_row - 3))
    with zipfile.ZipFile(path) as z:
        media = [n for n in z.namelist() if n.startswith("xl/media/")]
    print("- embedded chart images: %d" % len(media))
    ok &= len(media) == 4

    def cell(key, match, col):
        desc, cols, _ = d["sheets"][key]
        ws, hr = wb[names[key]], len(desc) + 1
        cj = [c.name for c in cols].index(col) + 1
        for i in range(hr + 1, ws.max_row + 1):
            if all(ws.cell(row=i, column=[c.name for c in cols].index(k) + 1).value == v for k, v in match.items()):
                return ws.cell(row=i, column=cj).coordinate, ws.cell(row=i, column=cj).value
        return None, None

    def doc_cell(fname, pred, j, conv=R.num):
        """A cell of the first results-doc table row matching pred(heading, table, row)."""
        for head, t in R.md_tables(os.path.join(RES, fname)):
            for c in t:
                if pred(head, t, c):
                    return conv(c[j])
        return None

    sk5 = d["sheets"]["s2"][2][1][1]
    spots = [
        ("s1", {"workload": "histo"}, "our IPC gain", "final_tier1.md §3, histo, IPC gain (ours)",
         doc_cell("final_tier1.md", lambda h, t, c: t[0][1].startswith("coverage: paper / pred") and
                  c[0].startswith("histo /"), 3, lambda s: R.nums(s)[1]), 1),
        ("s1", {"workload": "backprop"}, "perfect-L1 ceiling", "D1_ceilings.md §5 PAPER_V100, backprop, perfect L1 gain",
         doc_cell("D1_ceilings.md", lambda h, t, c: h.startswith("PAPER_V100") and c[0].startswith("backprop"), 3), 1),
        ("s2", {"config": dict(SK_LABELS)[BEST], "scope": sk5}, "accuracy_sc",
         "SK_match_qv100.md Means, SnakeCharmer apps, %s, accuracy_sc" % BEST,
         doc_cell("SK_match_qv100.md", lambda h, t, c: c[:2] == ["SnakeCharmer apps", BEST], 6), 1),
        ("s2", {"config": "SnakeCharmer SK (A100)"}, "on time", "SK_SC.md §4, SnakeCharmer SK (A100), on time",
         doc_cell("SK_SC.md", lambda h, t, c: t[0][0] == "metric" and c[0] == "on time", 1), 1),
        ("s4", {"workload": "particlefilter_naive"}, "IPC gain",
         "SK_match_qv100.md Per workload, particlefilter_naive, %s" % BEST,
         doc_cell("SK_match_qv100.md", lambda h, t, c: c[:2] == ["particlefilter_naive", BEST], 2), 2),
    ]
    print("\n| # | sheet | cell | workbook value | source value | source | match |")
    print("|---|---|---|---|---|---|---|")
    for i, (key, match, col, src, doc, dec) in enumerate(spots):
        coord, v = cell(key, match, col)
        good = v is not None and doc is not None and abs(round(v, dec) - doc) <= 0.5 * 10 ** -dec + 1e-9
        ok &= good
        print("| %d | %s | %s | %s | %s | %s | %s |" % (i + 1, names[key], coord, "n/a" if v is None else
                                                        ("%.*f" % (dec, v)), doc, src, "yes" if good else "NO"))
    return ok


def main():
    date = sys.argv[1] if len(sys.argv) > 1 else datetime.date.today().isoformat()
    outdir = os.path.join(RES, "slides")
    os.makedirs(outdir, exist_ok=True)
    d = collect()
    pngs = charts(d, date)
    for name, b in pngs.items():
        with open(os.path.join(outdir, name), "wb") as f:
            f.write(b)
    wb = Workbook()
    wb.remove(wb.active)
    add_sheets(wb, d, NAMES, pngs)
    out = os.path.join(outdir, "Snake_presentation_%s.xlsx" % date)
    wb.save(out)
    print(rel(out))
    for name in sorted(pngs):
        print(rel(os.path.join(outdir, name)))
    print("cross-checks: %d, mismatches: %d" % (len(CHECKS), len(MISMATCH)))
    for m in MISMATCH:
        print("  MISMATCH", m)
    ok = verify(out, d, NAMES)
    sys.exit(0 if ok and not MISMATCH else 1)


if __name__ == "__main__":
    main()
