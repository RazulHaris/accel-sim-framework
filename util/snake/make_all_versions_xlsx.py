#!/usr/bin/env python3
"""Every Snake version we have run, in one sheet
(docs/snake/results/Snake_all_versions_<date>.xlsx, sheet "Snake_results").

    PYTHONPATH=<dir with openpyxl> python3 util/snake/make_all_versions_xlsx.py [<date>]

No simulations.  Every number is read from existing outputs (sim_run_*
sim.out files) and results files; nothing is retyped.  A version's settings
are the effective option values of its runs (the option dump at the top of
each sim.out); options its build does not have are "not in this build".
A metric the build did not collect: "n/a: not collected in this build"; a
run that does not exist: "n/a: not run".
  Block 1  version definitions (one row per version)
  Block 2  results: one row per (version, conditions, workload), then summary
           rows per version and conditions; grouped PAPER_V100 1B, base
           QV100 100M, PAPER_V100 100M, PAPER_LAT 1B, references
  Block 3  metric definitions, warnings, caveats, doc-vs-options checks and
           provenance
The verification checks of the request are printed at the end (expected vs
got); exit status 1 if one fails.
"""
import datetime
import functools
import glob
import math
import os
import re
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
DOCS = os.path.join(ROOT, "docs", "snake")
from s4_report import parse  # noqa: E402
from sk_sc_report import metrics as sc_metrics, short, SEVEN, INPUT_DIFFERS, MISDECODED  # noqa: E402
from per_kernel import launches, deltas  # noqa: E402
import make_report_xlsx as R  # noqa: E402  (paper values, md tables)
from make_slides_data import sc_values, BRIEF_SC  # noqa: E402

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill, Alignment  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

P1, P, Q = "QV100-SASS-PAPER_V100-1B_INSN", "QV100-SASS-PAPER_V100-100M_INSN", "QV100-SASS-100M_INSN"
S = "-SNAKE-SNAKE_FP1"
SKS = S + "-SK_SC-SK_SC_CTA-SNAKE_HEAD64"
SK0 = SKS + "-SNAKE_MQR0-SNAKE_DM0"
X = SK0 + "-SNAKE_ACCF-SNAKE_ACCFB10K"
NAIVE = "particlefilter_naive-rodinia-3.1"
BFS = "bfs-rodinia-3.1"
NA_RUN, NA_BUILD = "n/a: not run", "n/a: not collected in this build"
NA_DIV = "n/a: nothing to divide"

# conditions: label, GPU setup, cap, baseline (dir, config) list, perfect-L1 (dir, config) or None
COND = {
    "P1": ("PAPER_V100", "1B", [("sim_run_baseline", P1)], ("sim_run_snake_gate/D1_pl1d", P1 + "-PERFECT_L1D")),
    "Q": ("base QV100", "100M", [("sim_run_snake_skc/qv100_7apps", Q), ("sim_run_snake_skc/qv100_sc", Q)], None),
    "P100": ("PAPER_V100", "100M", [("sim_run_snake_refcache/bin_ref_v2", P), ("sim_run_snake_skc/v100p_sc", P)], None),
    "LAT": ("PAPER_LAT", "1B", [("sim_run_baseline", P1 + "-PAPER_LAT")],
            ("sim_run_snake_gate/D1_pl1d", P1 + "-PAPER_LAT-PERFECT_L1D")),
}
ORDER = ["P1", "Q", "P100", "LAT"]
F = "sim_run_snake_fix/"
# name, kind, one-line description, {cond: [(run dir, config), ...]}, results doc
VERSIONS = [
    ("Baseline (no prefetching)", "base", "No prefetcher; the reference of every gain.",
     {"P1": [("sim_run_baseline", P1)], "Q": COND["Q"][2], "P100": COND["P100"][2]}, "results/final_tier1.md"),
    ("PAPER_LAT baseline", "base", "No prefetcher with the paper's measured L1 / L2 latencies (28 / 212 cycles).",
     {"LAT": [("sim_run_baseline", P1 + "-PAPER_LAT")]}, "results/final_tier1.md"),
    ("Tier 1 Snake", "snake", "Snake with the exact footprint: our paper-faithful default.",
     {"P1": [("sim_run_snake", P1 + S)],
      "Q": [("sim_run_snake_skc/qv100_7apps", Q + S), ("sim_run_snake_skc/qv100_sc", Q + S)],
      "P100": [("sim_run_snake_skc/v100p_7apps", P + S), ("sim_run_snake_skc/v100p_sc", P + S)]},
     "results/final_tier1.md"),
    ("Tier 1 Snake, line footprint", "snake", "Tier 1 Snake prefetching whole 128 B lines (footprint 0).",
     {"P1": [("sim_run_snake", P1 + "-SNAKE")]}, "results/final_tier1.md"),
    ("Snake-DT", "snake", "Paper ablation: no memory decoupling, no throttling.",
     {"P1": [("sim_run_snake_dt", P1 + "-SNAKE_DT-SNAKE_FP1")]}, "results/final_tier1.md"),
    ("Snake-T", "snake", "Paper ablation: decoupling without throttling.",
     {"P1": [("sim_run_snake_t", P1 + "-SNAKE_T-SNAKE_FP1")]}, "results/final_tier1.md"),
    ("s-Snake", "snake", "Paper comparison point: inter-thread chains only.",
     {"P1": [("sim_run_ssnake", P1 + "-S_SNAKE-SNAKE_FP1")]}, "results/final_tier1.md"),
    ("Snake-ISO", "snake", "Isolated-Snake: prefetches fill a separate buffer (decoupling mode 2).",
     {"P1": [("sim_run_snake", P1 + "-SNAKE_ISO-SNAKE_FP1")]}, "results/final_tier1.md"),
    ("PAPER_LAT Snake", "snake", "Tier 1 Snake with the paper's measured L1 / L2 latencies.",
     {"LAT": [("sim_run_snake", P1 + "-PAPER_LAT" + S)]}, "results/final_tier1.md"),
    ("SK_SC", "snake", "Snake with SnakeCharmer's published SK settings (cross-CTA inter-warp targets).",
     {"Q": [("sim_run_snake_skc/qv100_7apps", Q + S + "-SK_SC"), ("sim_run_snake_skc/qv100_sc", Q + S + "-SK_SC")],
      "P100": [("sim_run_snake_skc/v100p_7apps", P + S + "-SK_SC"), ("sim_run_snake_skc/v100p_sc", P + S + "-SK_SC")]},
     "results/SK_SC.md"),
    ("SK_SC_CTA", "snake", "SK_SC with inter-warp targets kept in the warp's CTA (hotspot and srad_v1 only).",
     {"Q": [("sim_run_snake_skc/cta", Q + S + "-SK_SC-SK_SC_CTA")],
      "P100": [("sim_run_snake_skc/cta", P + S + "-SK_SC-SK_SC_CTA")]}, "results/SK_SC.md"),
    ("SK", "snake", "SK match Step 1: SK_SC_CTA with SnakeCharmer's 64-entry per-warp Head.",
     {"Q": [("sim_run_snake_skm/qv100", Q + SKS)]}, "results/SK_match_qv100.md"),
    ("SK_R0", "snake", "SK with no miss-queue slots reserved for demand.",
     {"Q": [("sim_run_snake_skm/qv100", Q + SKS + "-SNAKE_MQR0")]}, "results/SK_match_qv100.md"),
    ("SK_R0_D0", "snake", "SK_R0 with memory decoupling off.",
     {"Q": [("sim_run_snake_skm/qv100", Q + SK0)], "P1": [("sim_run_snake_skm/v100p_1b", P1 + SK0)]},
     "results/SK_match_qv100.md"),
    ("A", "snake", "fix_today: Tier 1 Snake + R2 (demand first, MSHR at send).",
     {"P1": [(F + "p2", P1 + S + "-SNAKE_R2")]}, "results/fix_today.md"),
    ("B", "snake", "fix_today: SK_R0_D0 + R2 (demand first, MSHR at send).",
     {"P1": [(F + "p2", P1 + SK0 + "-SNAKE_R2")]}, "results/fix_today.md"),
    ("C", "snake", "fix_today: B with round-robin injection.",
     {"P1": [(F + "p2", P1 + SK0 + "-SNAKE_R2-SNAKE_PQARB1")]}, "results/fix_today.md"),
    ("D", "snake", "fix_today: Tier 1 Snake + accuracy filter (60%).",
     {"P1": [(F + "p3", P1 + S + "-SNAKE_ACCF")]}, "results/fix_today.md"),
    ("E", "snake", "fix_today: SK_R0_D0 + accuracy filter (60%).",
     {"P1": [(F + "p3", P1 + SK0 + "-SNAKE_ACCF")]}, "results/fix_today.md"),
    ("F", "snake", "fix_today: B + accuracy filter (60%).",
     {"P1": [(F + "p3b", P1 + SK0 + "-SNAKE_R2-SNAKE_ACCF")]}, "results/fix_today.md"),
    ("E10", "snake", "fix_today: E with a 10,000-cycle stop.",
     {"P1": [(F + "p3b", P1 + X)]}, "results/fix_today.md"),
    ("G8", "snake", "fix_today: E10 + one inter-warp target, the 8th resident warp ahead, CTA anchoring.",
     {"P1": [(F + "p4", P1 + X + "-SNAKE_IWD8")]}, "results/fix_today.md"),
    ("G16", "snake", "fix_today: E10 + one inter-warp target 16 ahead, CTA anchoring.",
     {"P1": [(F + "p4", P1 + X + "-SNAKE_IWD16")]}, "results/fix_today.md"),
    ("G32", "snake", "fix_today: E10 + one inter-warp target 32 ahead, CTA anchoring.",
     {"P1": [(F + "p4", P1 + X + "-SNAKE_IWD32")]}, "results/fix_today.md"),
    ("Perfect L1 (ceiling)", "ceil", "Every L1-cached global load hits at L1 latency (-gpgpu_perfect_l1d); no prefetcher.",
     {"P1": [("sim_run_snake_gate/D1_pl1d", P1 + "-PERFECT_L1D")],
      "LAT": [("sim_run_snake_gate/D1_pl1d", P1 + "-PAPER_LAT-PERFECT_L1D")]}, "results/D1_ceilings.md"),
    ("Perfect memory (ceiling)", "ceil", "Zero-cost memory beyond the SM (-gpgpu_perfect_mem); L1 misses still pay the L1 path.",
     {"P1": [("sim_run_snake_gate/D1_pmem", P1 + "-PERFECT_MEM")],
      "LAT": [("sim_run_snake_gate/D1_pmem", P1 + "-PAPER_LAT-PERFECT_MEM")]}, "results/D1_ceilings.md"),
]
SC_APPS = ["b+tree-rodinia-3.1", BFS, "particlefilter_float-rodinia-3.1", NAIVE,
           "heartwall-rodinia-2.0-ft", "streamcluster-rodinia-2.0-ft"]
APP_ORDER = list(SEVEN) + SC_APPS
VER, PREL, DIAG, EXCL, INDIFF, SUMM, REF = ("Verified (1B)", "Preliminary (100M)", "Diagnostic", "Excluded (bfs)",
                                           "Input differs", "Summary", "Reference")
FILL = {VER: "C6EFCE", PREL: "FFEB9C", DIAG: "D9D9D9", EXCL: "F8CBAD", INDIFF: "F8CBAD"}


# ------------------------------------------------------------ helpers
def sim(d, w, c):
    return os.path.join(ROOT, d, "runs", w[0], w[1], c, "sim.out")


def find(dirs_cfgs):
    out = {}
    for d, c in dirs_cfgs:
        for p in glob.glob(os.path.join(ROOT, d, "runs", "*", "*", c, "sim.out")):
            out[tuple(p.split(os.sep)[-4:-2])] = p
    return out


def wl_order(w):
    return (APP_ORDER.index(w[0]) if w[0] in APP_ORDER else 99, w[1])


@functools.lru_cache(maxsize=64)
def head(p):
    with open(p, errors="replace") as f:
        return f.read(400000)


def options(p):
    """Effective option values from a run's option dump."""
    o = {}
    for k, v in re.findall(r"(?m)^(-(?:snake|gpgpu)_\w+)\s+(.*?)\s*#", head(p)):
        o[k] = v.strip()
    return o


def tags(p):
    return " ".join(sorted(set(re.findall(r"(?:accelsim-commit|gpgpu-sim_git-commit)-[0-9a-f]+_modified_[0-9.]+",
                                          head(p)))))


@functools.lru_cache(maxsize=None)
def stats(p):
    return parse(p)


@functools.lru_cache(maxsize=None)
def launch_deltas(p):
    ls, capped = launches(p)
    return deltas(ls), capped


def best_kernel(bp, sp, cut):
    """(name, share of baseline cycles %, ratio) of the kernel with the largest IPC gain among kernels with
    >= 1% of the baseline cycles; launches pair in order; cut: drop the launch cut by the cap (1B)."""
    (db, cb), (ds, cs) = launch_deltas(bp), launch_deltas(sp)
    n = min(len(db), len(ds))
    if any(db[i][1] != ds[i][1] for i in range(n)):
        return None
    capped = cb or cs or len(db) != len(ds)
    idx = list(range(n - (1 if capped and cut else 0)))
    tot = sum(db[i][2] for i in idx)
    k = {}
    for i in idx:
        e = k.setdefault(db[i][1], [0, 0, 0, 0])
        e[0] += db[i][2]; e[1] += db[i][3]; e[2] += ds[i][2]; e[3] += ds[i][3]
    best = None
    for nm, (cb_, ib, cs_, is_) in k.items():
        if not cb_ or not cs_ or 100.0 * cb_ / tot < 1.0:
            continue
        r = (is_ / cs_) / (ib / cb_)
        if best is None or (r, cb_) > (best[2], best[3]):
            best = (nm, 100.0 * cb_ / tot, r, cb_)
    return best[:3] if best else None


def kname(m):
    g = re.match(r"_Z(\d+)", m)
    return m[g.end():g.end() + int(g.group(1))] if g else m


def pc(a, b):
    return 100.0 * a / b if b else None


def gm(xs):
    xs = [x for x in xs if isinstance(x, (int, float))]
    return (math.exp(sum(math.log1p(x / 100) for x in xs) / len(xs)) - 1) * 100 if xs else None


def am(xs):
    xs = [x for x in xs if isinstance(x, (int, float)) and x == x]
    return sum(xs) / len(xs) if xs else None


def same(a, b):
    """Option values equal as numbers or as strings."""
    try:
        return float(a) == float(b)
    except (TypeError, ValueError):
        return a == b


def git(path, *a):
    return subprocess.run(["git", "-C", path] + list(a), capture_output=True, text=True).stdout.strip()


# ------------------------------------------------------------ settings
def registered_defaults():
    """Option defaults as registered in the current gpgpu-sim source."""
    src = open(os.path.join(ROOT, "gpu-simulator", "gpgpu-sim", "src", "gpgpu-sim", "snake_prefetcher.cc")).read()
    return {m.group(1): m.group(2) for m in re.finditer(
        r'option_parser_register\(opp,\s*"(-snake_\w+)",.*?"([^"]*)"\);', src, re.S)}


def key_settings(o):
    g = lambda k: o.get("-snake_" + k, "not in this build")
    yn = lambda k: {"1": "on", "0": "off"}.get(g(k), g(k))
    scope = {"0": "same CTA (next warps of the CTA)", "1": "cross-CTA (resident warps by global warp id)",
             "2": "same warp index in the next CTAs"}.get(g("interwarp_scope"), g("interwarp_scope"))
    if g("interwarp_scope") == "not in this build":
        scope = "same CTA (scope option not in this build)"
    single = {"1": "single target (the degree-th)", "0": "all targets up to the degree"}.get(
        g("interwarp_single"), "all targets up to the degree (single option not in this build)")
    r2 = ("off (prefetches use the L1D miss queue)" if g("pq_path") in ("0", "not in this build") else
          "on: %s, MSHR %s, queue %s sectors" % ({"0": "demand first", "1": "round-robin"}[g("pq_arb")],
                                                 {"0": "at send", "1": "at enqueue"}[g("pq_mshr")], g("pq_size")))
    if g("pq_path") == "not in this build":
        r2 = "off (R2 not in this build)"
    acc = ("off" if g("entry_acc_filter") in ("0", "not in this build") else
           "on: threshold %s, window %s resolved, unused after %s cycles, stop %s cycles" % (
               g("acc_filter_min"), g("acc_filter_window"), g("acc_filter_age"), g("acc_filter_block")))
    if g("entry_acc_filter") == "not in this build":
        acc = "off (filter not in this build)"
    deep = {"1": "promoted entry", "0": "trained entry"}.get(g("chain_deep_promoted"),
                                                              "trained entry (option not in this build)")
    intra = ("%s, %s" % (g("intra_distance"), {"0": "every iteration 1..D", "1": "only iteration D"}.get(
        g("intra_distance_mode"), g("intra_distance_mode")))) if g("intra_distance") != "not in this build" \
        else "1, next iteration (option not in this build)"
    return "; ".join([
        "sources: inter-thread %s, inter-warp %s, intra-warp %s" % (yn("enable_inter_thread"), yn("enable_inter_warp"),
                                                                      yn("enable_intra_warp")),
        "footprint %s" % {"0": "line (0)", "1": "exact (1)"}.get(g("footprint"), g("footprint")),
        "lane filter: lane_mode %s, require_equal_lane_stride %s" % (g("lane_mode"), g("require_equal_lane_stride")),
        "Head: %s, %s rows" % ({"1": "two warps per PC (head_two_warps 1)",
                                "0": "64-entry per-warp (head_two_warps 0)"}.get(g("head_two_warps"), g("head_two_warps")),
                               g("head_entries")),
        "Tail: %s entries, %s" % (g("tail_entries"), g("tail_evict")),
        "training: promote %s warps, intra confirm %s, inter-warp min %s, promoted_all_warps %s, demote_removed %s" % (
            g("promote_warps"), g("intra_confirm_warps"), g("interwarp_min_warps"), g("promoted_all_warps"),
            g("demote_removed")),
        "chain depth %s (deeper steps: %s; while throttled %s)" % (g("max_chain_depth"), deep, g("throttle_chain_depth")),
        "intra-warp distance %s" % intra,
        "inter-warp: degree %s, %s, %s, CTA anchoring %s, chain from targets %s" % (
            g("interwarp_degree"), scope, single,
            {"1": "on", "0": "off"}.get(g("interwarp_cta_anchor"), "off (not in this build)"), g("interwarp_chain_depth")),
        "prefetch queue %s lines, issue %s/cycle, latency %s" % (g("queue_size"), g("issue_width"), g("latency")),
        "miss-queue reserve %s, MSHR reserve %s" % (g("missq_reserve"), g("mshr_reserve")),
        "R2: %s" % r2,
        "accuracy filter: %s" % acc,
        "decoupling: mode %s (%s), variant %s" % (g("decouple_mode"), {"0": "none", "1": "unified region",
                                                                       "2": "isolated buffer"}.get(g("decouple_mode"), "?"),
                                                 g("decouple_variant")),
        "throttling: %s (bandwidth meter mode %s, halt >= %s, resume <= %s, window %s cycles; space halt %s cycles)" % (
            yn("throttle_enable"), g("bw_mode"), g("bw_high"), g("bw_low"), g("bw_window"), g("throttle_space_cycles")),
    ])


def cond_settings(o):
    clk = o.get("-gpgpu_clock_domains", "?").split(":")[0]
    extra = []
    for k in ("-gpgpu_l1_latency", "-gpgpu_l2_rop_latency"):
        extra.append("%s %s" % (k[1:], o.get(k, "?")))
    for k in ("-gpgpu_perfect_l1d", "-gpgpu_perfect_mem"):
        if o.get(k) == "1":
            extra.append("%s 1" % k[1:])
    return "scheduler %s, core clock %s MHz, -gpgpu_max_insn %s, %s" % (
        o.get("-gpgpu_scheduler", "?"), clk, o.get("-gpgpu_max_insn", "?"), ", ".join(extra))


# origin of options that differ from Tier 1 Snake: (origin, OQ)
ORIGIN = {
    "-snake_footprint": ("prefetch footprint: line (Tier 1 uses exact, our reading)", "OQ-9"),
    "-snake_queue_size": ("SnakeCharmer setting", "OQ-33"),
    "-snake_max_chain_depth": ("SnakeCharmer setting", "OQ-7, OQ-33"),
    "-snake_chain_deep_promoted": ("SnakeCharmer setting", "OQ-35"),
    "-snake_intra_distance": ("SnakeCharmer setting", "OQ-34"),
    "-snake_intra_distance_mode": ("SnakeCharmer setting", "OQ-34"),
    "-snake_head_two_warps": ("SnakeCharmer's 64-entry per-warp Head", "OQ-1b, OQ-37"),
    "-snake_missq_reserve": ("our implementation parameter", "OQ-11"),
    "-snake_decouple_mode": ("decoupling (paper §3.2) changed", "OQ-12"),
    "-snake_throttle_enable": ("throttling (paper §3.3) off: paper ablation", "§4"),
    "-snake_enable_inter_warp": ("prediction source off: paper comparison point", "§4"),
    "-snake_enable_intra_warp": ("prediction source off: paper comparison point", "§4"),
    "-snake_pq_path": ("R2: our reading of SnakeCharmer §III-A", "OQ-38"),
    "-snake_pq_arb": ("R2", "OQ-38"), "-snake_pq_mshr": ("R2", "OQ-38"), "-snake_pq_size": ("R2", "OQ-38"),
    "-snake_entry_acc_filter": ("our extension", "OQ-39"), "-snake_acc_filter_min": ("our extension", "OQ-39"),
    "-snake_acc_filter_block": ("our extension", "OQ-39"), "-snake_acc_filter_window": ("our extension", "OQ-39"),
    "-snake_acc_filter_age": ("our extension", "OQ-39"),
    "-snake_interwarp_single": ("our extension", "OQ-40"), "-snake_interwarp_cta_anchor": ("our extension", "OQ-40"),
    "-gpgpu_l1_latency": ("the paper's measured L1 latency", "SPEC §5"),
    "-gpgpu_l2_rop_latency": ("the paper's measured L2 latency", "SPEC §5"),
    "-gpgpu_perfect_l1d": ("ceiling study, not a prefetcher", "D1"), "-gpgpu_perfect_mem": ("ceiling study", "D1"),
}
TIER1_TEXT = ("Paper-faithful default (Snake, MICRO'23 §3). Our readings where the paper is open: exact footprint "
              "(OQ-9), lane filter mode 1 (OQ-17), decoupling mode 1 variant 1 (OQ-12, OQ-15), bandwidth meter on "
              "miss-queue sends (OQ-10), miss-queue reserve 8 and MSHR reserve 8 (OQ-11), inter-warp targets: next "
              "4 warps of the CTA (OQ-8), Tail 10 lru_popcount (Figs. 20, 22).")


def origin_text(o, ref, defaults):
    diffs = []
    for k in sorted(set(o) | set(ref)):
        if not (k.startswith("-snake_") or k in ORIGIN):
            continue
        if k in ("-snake_enable", "-snake_debug_trace", "-snake_energy_pj_per_access", "-snake_static_mw"):
            continue
        a = o.get(k)
        b = ref.get(k, defaults.get(k))
        if a is None or b is None or a == b:
            continue
        what, oq = ORIGIN.get(k, ("changed", ""))
        if k == "-snake_interwarp_degree" or k == "-snake_interwarp_scope":
            what, oq = (("our extension (single target / anchoring)", "OQ-40") if o.get("-snake_interwarp_single") == "1"
                        else ("SnakeCharmer setting (scope 1 from the S6 experiment)", "OQ-33") if
                        o.get("-snake_interwarp_scope") == "1" else ("SnakeCharmer setting, same-CTA targets", "OQ-33, OQ-37"))
        diffs.append("%s %s (Tier 1 %s): %s%s" % (k[1:], a, b, what, (", " + oq) if oq else ""))
    return diffs


# ------------------------------------------------------------ results
def inputs_note():
    """Input match per SnakeCharmer app, from SK_SC.md §1."""
    out = {}
    for cells in sum([t for _, t in R.md_tables(os.path.join(DOCS, "results", "SK_SC.md"))], []):
        if len(cells) >= 3 and cells[2] in ("same", "**differs**"):
            out[cells[0].split(" (")[0]] = "yes (SnakeCharmer Table II)" if cells[2] == "same" else \
                "no (differs from SnakeCharmer Table II)"
    return out


def run_row(vname, kind, ck, w, sp, bp, ceil, paper, innote):
    gpu, cap = COND[ck][0], COND[ck][1]
    a = w[0]
    status = EXCL if a == BFS else INDIFF if a.startswith(INPUT_DIFFERS) else (VER if cap == "1B" else PREL)
    inp = "Snake paper gives no inputs (SPEC OQ-13)" if a in SEVEN else innote.get(
        a.replace("-rodinia-3.1", "").replace("-rodinia-2.0-ft", "").split("_")[0], "?")
    s = stats(sp)
    b = stats(bp) if bp else None
    r = {"Version": vname, "GPU setup": gpu, "Instruction cap": cap, "Workload": short(*w), "Input matches paper?": inp,
         "Run dir": os.path.relpath(os.path.dirname(sp), ROOT), "Status": status,
         "IPC (whole run)": s["_ipc"], "_build": tags(sp), "_w": w, "_ok": s["_done"]}
    r["IPC gain vs no prefetching (%)"] = (s["_ipc"] / b["_ipc"] - 1) * 100 if b and kind != "base" else (
        "n/a: this is the baseline" if kind == "base" else NA_RUN)
    if kind != "base" and bp:
        bk = best_kernel(bp, sp, cap == "1B")
        r["Best kernel (Diagnostic)"] = kname(bk[0]) if bk else "n/a: no kernel >= 1%"
        r["Best kernel share of run time (%)"] = bk[1] if bk else None
        r["Best kernel IPC / same kernel without prefetching (Diagnostic)"] = bk[2] if bk else None
    else:
        for k in ("Best kernel (Diagnostic)", "Best kernel share of run time (%)",
                  "Best kernel IPC / same kernel without prefetching (Diagnostic)"):
            r[k] = "n/a: this is the baseline"
    r["Perfect-L1 ceiling (%)"] = ceil.get(w, NA_RUN) if ceil is not None else NA_RUN
    snake = "demand_sector_accesses" in s and kind == "snake"
    none = "n/a: no prefetcher" if kind != "snake" else NA_BUILD
    dem = s.get("demand_sector_accesses", 0)
    if snake:
        tl, lt, iss = s.get("pf_useful_timely", 0), s.get("pf_late", 0), s.get("pf_issued", 0)
        r["Coverage, Snake-paper style (coverage_pred, %)"] = pc(s.get("demand_pred_matched", 0), dem)
        r["Delivered coverage (timely + late) / demand (%)"] = pc(tl + lt, dem)
        r["Accuracy, Snake-paper style (timely / demand, %)"] = pc(tl, dem)
        r["Prefetches issued per demand sector"] = iss / dem if dem else NA_DIV
        r["Unused share of issued (%)"] = pc(iss - tl - lt, iss) if iss else NA_DIV
    else:
        for k in ("Coverage, Snake-paper style (coverage_pred, %)", "Delivered coverage (timely + late) / demand (%)",
                  "Accuracy, Snake-paper style (timely / demand, %)", "Prefetches issued per demand sector",
                  "Unused share of issued (%)"):
            r[k] = none
    sc_cols = ["SnakeCharmer coverage H/(H+M), H = on-time (%)", "SnakeCharmer coverage, H = on-time + late (%)",
               "SnakeCharmer accuracy H/F (%)", "On time (%)", "Too early (%)", "Too late (%)",
               "Prefetch-induced demand evictions / all victim sectors (%)"]
    if snake and "sc_pf_fills" in s:
        m = sc_metrics(s)
        for col, k in zip(sc_cols, ("cov_on", "cov_late", "acc", "t_on", "t_early", "t_late", "pfev")):
            r[col] = m[k] if m[k] == m[k] else NA_DIV
    else:
        for col in sc_cols:
            r[col] = none
    r["L1 hit rate, baseline (%)"] = pc(b["_hit"], b["_acc"]) if b else NA_RUN
    r["L1 hit rate, version (%)"] = pc(s["_hit"], s["_acc"])
    pv = paper.get(a.replace("-rodinia-3.1", "").replace("parboil-", ""))
    for col, i in (("Snake paper coverage (%)", 0), ("Snake paper accuracy (%)", 1), ("Snake paper IPC gain (%)", 2)):
        r[col] = pv[i] if pv else "n/a: not a Snake paper app"
    return r


COLS = ["Version", "GPU setup", "Instruction cap", "Workload", "Input matches paper?", "Run dir", "Status",
        "IPC (whole run)", "IPC gain vs no prefetching (%)", "Best kernel (Diagnostic)",
        "Best kernel share of run time (%)", "Best kernel IPC / same kernel without prefetching (Diagnostic)",
        "Perfect-L1 ceiling (%)", "Coverage, Snake-paper style (coverage_pred, %)",
        "Delivered coverage (timely + late) / demand (%)", "Accuracy, Snake-paper style (timely / demand, %)",
        "SnakeCharmer coverage H/(H+M), H = on-time (%)", "SnakeCharmer coverage, H = on-time + late (%)",
        "SnakeCharmer accuracy H/F (%)", "On time (%)", "Too early (%)", "Too late (%)",
        "Prefetch-induced demand evictions / all victim sectors (%)", "Prefetches issued per demand sector",
        "Unused share of issued (%)", "L1 hit rate, baseline (%)", "L1 hit rate, version (%)",
        "Snake paper coverage (%)", "Snake paper accuracy (%)", "Snake paper IPC gain (%)", "Build"]
MEANCOLS = ["Coverage, Snake-paper style (coverage_pred, %)", "Delivered coverage (timely + late) / demand (%)",
            "Accuracy, Snake-paper style (timely / demand, %)", "SnakeCharmer coverage H/(H+M), H = on-time (%)",
            "SnakeCharmer coverage, H = on-time + late (%)", "SnakeCharmer accuracy H/F (%)", "On time (%)",
            "Too early (%)", "Too late (%)", "Prefetch-induced demand evictions / all victim sectors (%)",
            "Prefetches issued per demand sector", "Unused share of issued (%)", "L1 hit rate, baseline (%)",
            "L1 hit rate, version (%)"]


def summaries(vname, kind, ck, rows):
    """Summary rows for one version and conditions."""
    out = []
    gpu, cap = COND[ck][0], COND[ck][1]
    byw = {r["_w"]: r for r in rows}
    seven = [w for w in byw if w[0] in SEVEN]
    sc = [w for w in byw if w[0] in SC_APPS and w[0] != BFS]
    sets = []
    if len(seven) == 9:
        sets.append(("our 7 apps (an app's inputs combined by GMEAN; rates: mean over apps, inputs first)", seven, True))
        sets.append(("our 9 workloads", seven, False))
    if sc:
        sets.append(("SnakeCharmer's workloads (%d, without bfs)" % len(sc), sc, False))
        if any(w[0] == NAIVE for w in sc):
            sets.append(("SnakeCharmer's workloads without particlefilter_naive (%d)" % (len(sc) - 1),
                         [w for w in sc if w[0] != NAIVE], False))
        if any(w[0] == BFS for w in byw):
            sets.append(("SnakeCharmer's workloads (%d, with bfs, as SK_SC.md §4)" % (len(sc) + 1),
                         sc + [w for w in byw if w[0] == BFS], False))
    if len(seven) == 9 and sc:
        allw = seven + sc
        sets.append(("all %d workloads (without bfs)" % len(allw), allw, False))
        sets.append(("all workloads without particlefilter_naive (%d)" % (len(allw) - 1),
                     [w for w in allw if w[0] != NAIVE], False))
    for label, ws, apps in sets:
        def agg(col, geo):
            if apps:
                per = {}
                for w in ws:
                    per.setdefault(w[0], []).append(byw[w][col])
                vals = [(gm(v) if geo else am(v)) for v in per.values()]
                return gm(vals) if geo else am(vals)
            return gm([byw[w][col] for w in ws]) if geo else am([byw[w][col] for w in ws])
        r = {c: "" for c in COLS}
        r.update({"Version": vname, "GPU setup": gpu, "Instruction cap": cap, "Workload": "GMEAN / mean: " + label,
                  "Status": SUMM, "Run dir": "", "Input matches paper?": ""})
        if kind != "base":
            r["IPC gain vs no prefetching (%)"] = agg("IPC gain vs no prefetching (%)", True)
            pfl = [byw[w]["Perfect-L1 ceiling (%)"] for w in ws]
            r["Perfect-L1 ceiling (%)"] = agg("Perfect-L1 ceiling (%)", True) if all(
                isinstance(x, float) for x in pfl) else NA_RUN
        for c in MEANCOLS:
            v = agg(c, False)
            r[c] = v if v is not None else (byw[ws[0]][c] if isinstance(byw[ws[0]][c], str) else NA_DIV)
        out.append(r)
        if kind != "base":
            d = {c: "" for c in COLS}
            d.update({"Version": vname, "GPU setup": gpu, "Instruction cap": cap, "Status": DIAG,
                      "Workload": "GMEAN best-kernel ratio (Diagnostic): " + label})
            g = agg("_bkgain", True)
            d["Best kernel IPC / same kernel without prefetching (Diagnostic)"] = 1 + g / 100 if g is not None else NA_RUN
            out.append(d)
    return out


# ------------------------------------------------------------ main
def main():
    date = sys.argv[1] if len(sys.argv) > 1 else datetime.date.today().isoformat()
    paper, prow, _ = R.paper_values()
    innote = inputs_note()
    defaults = registered_defaults()
    base = {ck: find(COND[ck][2]) for ck in COND}
    ceil = {}
    for ck, c in COND.items():
        if c[3]:
            cp = find([c[3]])
            ceil[ck] = {w: (stats(p)["_ipc"] / stats(base[ck][w])["_ipc"] - 1) * 100 for w, p in cp.items() if w in base[ck]}
        else:
            ceil[ck] = None
    results, vdef, prov, mism, bad = {}, [], [], [], []
    t1_opts = options(sim("sim_run_snake", ("parboil-histo", "_i___data_default_input_img_bin__o_ref_bmp____20_4"),
                          P1 + S))
    for vname, kind, desc, runs, doc in VERSIONS:
        conds, builds, eff = [], {}, None
        for ck in ORDER:
            if ck not in runs:
                continue
            paths = find(runs[ck])
            if not paths:
                continue
            rows = []
            for w in sorted(paths, key=wl_order):
                sp = paths[w]
                bp = base[ck].get(w) if kind != "base" else sp
                r = run_row(vname, kind, ck, w, sp, bp, ceil[ck], paper, innote)
                r["_bkgain"] = (r["Best kernel IPC / same kernel without prefetching (Diagnostic)"] - 1) * 100 if isinstance(
                    r["Best kernel IPC / same kernel without prefetching (Diagnostic)"], float) else None
                r["Build"] = r["_build"]
                if not r["_ok"]:
                    bad.append("%s %s %s: run did not finish" % (vname, ck, short(*w)))
                rows.append(r)
                builds[r["_build"]] = builds.get(r["_build"], 0) + 1
            results[(vname, ck)] = rows
            o = options(paths[sorted(paths, key=wl_order)[0]])
            if eff is None:
                eff = o
            # the same version under other conditions (possibly another build): options both builds have
            # must agree
            if kind == "snake":
                diff = sorted(k for k in set(o) & set(eff) if k.startswith("-snake_") and not same(o[k], eff[k]))
                if diff:
                    mism.append((vname, "Snake options differ between conditions: " + ", ".join(
                        "%s %s vs %s" % (k[1:], eff[k], o[k]) for k in diff)))
            conds.append("%s %s: %s; runs %s (%d workloads)" % (COND[ck][0], COND[ck][1], cond_settings(o),
                                                              ", ".join(sorted({d for d, _ in runs[ck]})), len(rows)))
            # tokens vs effective
            cfg = runs[ck][0][1]
            yml = open(os.path.join(ROOT, "util", "job_launching", "configs", "define-standard-cfgs.yml")).read()
            for tok in cfg.split("-")[2:]:
                m = re.search(r"(?m)^%s:\n\s+extra_params: \"((?:[^\"])*)\"" % re.escape(tok), yml)
                if not m:
                    continue
                for k, v in re.findall(r"(-(?:snake|gpgpu)_\w+)\s+([^\s\\\\]+)", m.group(1)):
                    if o.get(k) is not None and not same(o.get(k), v) and not (k == "-gpgpu_clock_domains"):
                        mism.append((vname, "token %s sets %s %s, effective %s (a later token overrides it)" % (
                            tok, k, v, o.get(k))))
        if eff is None:
            vdef.append([vname, desc, NA_RUN, "", NA_RUN, ""])
            continue
        if kind == "base":
            ks, org = "no prefetcher (-snake_enable %s)" % eff.get("-snake_enable", "not in this build"), "Reference"
        elif kind == "ceil":
            ks = "no prefetcher; " + ("-gpgpu_perfect_l1d 1" if eff.get("-gpgpu_perfect_l1d") == "1" else
                                      "-gpgpu_perfect_mem 1")
            org = "Ceiling study (D1), not a prefetcher"
        else:
            ks = key_settings(eff)
            org = TIER1_TEXT if vname == "Tier 1 Snake" else "Differs from Tier 1 Snake in: " + (
                "; ".join(origin_text(eff, t1_opts, defaults)) or "nothing (same options; other conditions)")
        if vname.startswith("PAPER_LAT"):
            org = (org + " " if kind == "snake" else "") + "Latencies: -gpgpu_l1_latency %s, -gpgpu_l2_rop_latency %s " \
                "(the paper's measured 28 / 212 cycles, SPEC §5)." % (eff.get("-gpgpu_l1_latency"),
                                                                       eff.get("-gpgpu_l2_rop_latency"))
        vdef.append([vname, desc, ks, org, "\n".join(conds),
                     "; ".join("%s (%d runs)" % (b, n) for b, n in sorted(builds.items()))])
        prov.append([vname, "; ".join("%s (%d runs)" % (b, n) for b, n in sorted(builds.items())),
                     doc, git(ROOT, "log", "-1", "--format=%h %ad", "--date=short", "--", "docs/snake/" + doc)])
    # reference versions
    sc = sc_values()
    pavg = {nm: v for nm, v, _ in prow}.get("Avg / GMEAN")
    vdef.append(["Snake paper (MICRO'23)", "Mostofi et al.: published per-app values (Figs. 16–18), read off bar charts "
                 "(±3 pp; SPEC §5).", "V100 model in Accel-Sim (80 SMs, 1530 MHz, GTO), to completion or 1B "
                 "instructions (SPEC §5)", "The paper", "Their simulator and inputs (not given, OQ-13)", "SPEC §5"])
    vdef.append(["SnakeCharmer SK (MICRO'26)", "Tang et al.: published averages of their Snake baseline SK "
                 "(Fig. 12 AVG, Fig. 4).", "A100 (1410 MHz, 108 SMs), zero-compressed cache; SK settings as in "
                 "SK_SC (OQ-33)", "The paper", "A100, 14 apps, full runs", "results/SK_SC.md §4"])

    # ---------------------------------------------------------- workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Snake_results"
    bold = Font(bold=True)
    row = 1
    ws.cell(row=row, column=1, value="Snake: every version we ran, its effective settings and all results "
            "(generated %s by util/snake/make_all_versions_xlsx.py; no new simulations)" % date).font = Font(bold=True, size=13)
    row += 2
    ws.cell(row=row, column=1, value="BLOCK 1: VERSION DEFINITIONS").font = Font(bold=True, size=12)
    row += 1
    b1h = ["Version", "One-line description", "Key settings (effective values from the runs' option dumps)",
           "Paper-faithful or our extension (which parts, OQ numbers)", "Conditions it was run under",
           "Commit / build tag"]
    for j, h in enumerate(b1h):
        ws.cell(row=row, column=j + 1, value=h).font = bold
    row += 1
    for v in vdef:
        for j, x in enumerate(v):
            c = ws.cell(row=row, column=j + 1, value=x)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="BLOCK 2: RESULTS (one row per version, conditions and workload; summary rows "
            "after each version)").font = Font(bold=True, size=12)
    row += 1
    h2 = row
    for j, h in enumerate(COLS):
        c = ws.cell(row=row, column=j + 1, value=h)
        c.font = bold
        c.alignment = Alignment(wrap_text=True, vertical="top")
    row += 1
    pct_cols = {c for c in COLS if c.endswith("(%)") or c.endswith("%)")}
    ratio_cols = {"Best kernel IPC / same kernel without prefetching (Diagnostic)"}
    gain_cols = {"IPC gain vs no prefetching (%)", "Perfect-L1 ceiling (%)", "Snake paper IPC gain (%)"}
    all_rows = []
    for ck in ORDER:
        for vname, kind, _, _, _ in VERSIONS:
            rows = results.get((vname, ck))
            if not rows:
                continue
            all_rows += rows + summaries(vname, kind, ck, rows)
    # references
    ref = {c: "" for c in COLS}
    ref.update({"Version": "Snake paper (MICRO'23)", "GPU setup": "V100 (paper)", "Instruction cap": "full or 1B",
                "Workload": "mean over our 7 apps (IPC: GMEAN)", "Status": REF,
                "Snake paper coverage (%)": am([paper[a][0] for a in R.APPS]),
                "Snake paper accuracy (%)": am([paper[a][1] for a in R.APPS]),
                "Snake paper IPC gain (%)": gm([paper[a][2] for a in R.APPS])})
    ref2 = dict(ref, Workload="the paper's own average over its 11 apps (table row 'Avg / GMEAN')",
                **{"Snake paper coverage (%)": pavg[0], "Snake paper accuracy (%)": pavg[1],
                   "Snake paper IPC gain (%)": pavg[2]})
    scr = {c: "" for c in COLS}
    scr.update({"Version": "SnakeCharmer SK (MICRO'26)", "GPU setup": "A100 (theirs)", "Instruction cap": "full runs",
                "Workload": "average over their 14 apps (SK_SC.md §4; evictions from the fix_today brief)",
                "Status": REF, "SnakeCharmer coverage H/(H+M), H = on-time (%)": sc.get("cov"),
                "SnakeCharmer coverage, H = on-time + late (%)": "n/a: SnakeCharmer reports one coverage",
                "SnakeCharmer accuracy H/F (%)": sc.get("acc"), "On time (%)": sc.get("on"),
                "Too early (%)": sc.get("early"), "Too late (%)": sc.get("late"),
                "Prefetch-induced demand evictions / all victim sectors (%)": BRIEF_SC["evict"],
                "IPC gain vs no prefetching (%)": "n/a: not reported for SK"})
    all_rows += [ref, ref2, scr]
    for r in all_rows:
        for j, col in enumerate(COLS):
            v = r.get(col, "")
            if v is None:
                v = NA_DIV
            c = ws.cell(row=row, column=j + 1, value=v)
            if isinstance(v, float):
                c.number_format = "0.0000" if col in ratio_cols else (
                    '+0.00"%";-0.00"%";0.00"%"' if col in gain_cols else ('0.00"%"' if col in pct_cols else "0.00"))
            if col == "Status" and v in FILL:
                c.fill = PatternFill("solid", fgColor=FILL[v])
        if r.get("Status") in (SUMM, DIAG, REF):
            for j in range(4):
                ws.cell(row=row, column=j + 1).font = Font(italic=True)
        row += 1
    last2 = row - 1
    ws.auto_filter.ref = "A%d:%s%d" % (h2, get_column_letter(len(COLS)), last2)
    row += 1
    # ---------------------------------------------------------- block 3
    ws.cell(row=row, column=1, value="BLOCK 3: METRIC DEFINITIONS AND NOTES").font = Font(bold=True, size=12)
    row += 1
    defs = [
        ("IPC gain vs no prefetching", "gpu_tot_ipc(version) / gpu_tot_ipc(no prefetching, same GPU setup, cap and "
         "workload) − 1", "sim.out"),
        ("Best kernel (Diagnostic)", "per workload, the kernel function with the largest IPC ratio among kernels with "
         ">= 1% of the no-prefetch run's cycles; kernel IPC = instructions / cycles over its launches (paired in "
         "order). At 1B the launch cut by the cap is excluded; at 100M it is kept (it is the only launch of backprop "
         "and hotspot 1024). Biased upward by selection: diagnostic only.", "K1_best_kernel.md, kernel_gap.py"),
        ("Perfect-L1 ceiling", "IPC with -gpgpu_perfect_l1d 1 / no-prefetch IPC − 1 (every L1-cached global load "
         "hits); run only at PAPER_V100 and PAPER_LAT, 1B", "D1_ceilings.md"),
        ("Coverage, Snake-paper style (coverage_pred)", "demand sectors an earlier prediction named (issued or not, "
         "within 5,000 cycles) / demand sectors (GLOBAL_ACC_R accesses not failing with RESERVATION_FAIL)",
         "Snake paper §4 ('correctly predicted addresses / total demand addresses'); SPEC OQ-14"),
        ("Delivered coverage", "(timely + late prefetched sectors) / demand sectors (coverage_paper)", "SPEC OQ-14"),
        ("Accuracy, Snake-paper style (accuracy_paper)", "timely prefetched sectors / demand sectors: timely "
         "coverage, not useful / issued", "Snake paper §4 ('timely correctly predicted addresses'); SPEC OQ-14"),
        ("SnakeCharmer coverage H/(H+M)", "H = loads hitting a prefetched sector (on-time) [+ late merges in the second "
         "column]; M = loads sending a new request (late counted in M in the first column)",
         "SnakeCharmer §II-B, Fig. 12 caption; SPEC OQ-36"),
        ("SnakeCharmer accuracy H/F", "(on-time uses + fills a load had merged into) / prefetch fills "
         "((sc_ld_ontime + sc_pf_fills − pf_filled) / sc_pf_fills)", "SnakeCharmer §II-B; SPEC OQ-36"),
        ("On time / too early / too late", "shares of consumed prefetches (first load): hit after the fill / evicted "
         "unused earlier in the kernel / merged into the in-flight MSHR entry", "SnakeCharmer §II-B, Fig. 4; SPEC OQ-36"),
        ("Prefetch-induced demand evictions", "demand-filled victim sectors evicted by a prefetch allocation / all "
         "valid victim sectors", "SnakeCharmer §II-B; SPEC OQ-36"),
        ("Prefetches issued per demand sector", "snake_pf_issued / snake_demand_sector_accesses", "our stats"),
        ("Unused share of issued", "(issued − timely − late) / issued: evicted unused plus never used by the end",
         "our stats (SNAKE_PLAN §7 outcome invariant)"),
        ("L1 hit rate", "Total_core_cache_stats_breakdown[GLOBAL_ACC_R][HIT] / [TOTAL_ACCESS]", "sim.out"),
        ("Summary rows", "IPC: GMEAN of gains; '7 apps' combines an app's inputs by GMEAN first (hotspot, lud). "
         "Coverage, accuracy and the other rates: unweighted arithmetic mean over the workloads that have a value "
         "(over apps, inputs first, in the '7 apps' row); a workload with nothing to divide (e.g. nw, no prefetch "
         "used) is skipped", "this script"),
        ("WARNING", "The Snake paper's 'accuracy' (timely / demand) and SnakeCharmer's (consumed / fills) are "
         "different quantities. Never compare them directly.", ""),
        ("Caveat: traces", "All traces are sm_70 (V100, hw_run). No real A100 traces exist locally (SK_SC.md §2).", ""),
        ("Caveat: SnakeCharmer values", "A100 averages over 14 apps (only 5 of ours, two with other inputs), full runs, "
         "with a zero-compressed cache. Not reproducible here.", "SK_SC.md §4"),
        ("Caveat: best kernel", "Diagnostic and biased upward (the maximum over kernels). It does not predict the "
         "whole run (histo's best kernel gains while every histo run loses).", "fix_today.md"),
        ("Caveat: particlefilter_naive", "Its +65–79% dominates plain averages over SnakeCharmer's workloads; see the "
         "rows without it.", "SK_SC.md §5"),
        ("Caveat: conditions", "100M vs 1B and PAPER_V100 (GTO, 1530 MHz) vs base QV100 (LRR, 1132 MHz) must not be "
         "compared directly (window and setup effects: SK_match_qv100.md, decomposition section).", ""),
        ("Not included", "Stage gate runs (S2–S6, D1 off-regressions), Tier 1 parameter sweeps (Tail, eviction policy, "
         "throttle halt) and the Tier 2 sensitivity runs: in Snake_results_<date>.xlsx and the stage results files.",
         ""),
        ("Status colours", "Verified (1B) green; Preliminary (100M) yellow; Diagnostic grey; Excluded (bfs: 24.8% of "
         "its global memory lines mis-decoded) and Input differs (heartwall, streamcluster) orange", ""),
    ]
    for j, h in enumerate(["Item", "Definition / formula / note", "Source"]):
        ws.cell(row=row, column=j + 1, value=h).font = bold
    row += 1
    for d in defs:
        for j, x in enumerate(d):
            ws.cell(row=row, column=j + 1, value=x).alignment = Alignment(wrap_text=True, vertical="top")
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="Settings checks (tokens and docs vs effective options)").font = bold
    row += 1
    # docs vs registered defaults: SNAKE_PLAN options table
    for t in R.md_tables_text(open(os.path.join(DOCS, "SNAKE_PLAN.md")).read()):
        for cells in t:
            m = re.match(r"`(snake_\w+)`", cells[0]) if cells else None
            if not m or len(cells) < 3:
                continue
            k = "-" + m.group(1)
            doc_v = cells[2].strip("`").split(" ")[0].split("/")[0]  # "1/1/1": one row for several options
            src_v = defaults.get(k)
            if src_v is not None and doc_v not in ("", "-") and doc_v != src_v and \
                    not (re.fullmatch(r"[\d.]+", doc_v) and re.fullmatch(r"[\d.]+", src_v) and float(doc_v) == float(src_v)):
                mism.append(("SNAKE_PLAN options table", "%s: doc default %s, source default %s (the effective "
                             "options are used)" % (k[1:], doc_v, src_v)))
    for v, m_ in (mism or [("all", "no mismatch found")]):
        ws.cell(row=row, column=1, value=v)
        ws.cell(row=row, column=2, value=m_).alignment = Alignment(wrap_text=True)
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="Provenance").font = bold
    row += 1
    for j, h in enumerate(["Version", "Build tags (runs)", "Results file", "Results file commit"]):
        ws.cell(row=row, column=j + 1, value=h).font = bold
    row += 1
    for p_ in prov + [["trees when generated", "accel-sim %s, gpgpu-sim %s" % (
            git(ROOT, "log", "-1", "--format=%h"), git(os.path.join(ROOT, "gpu-simulator", "gpgpu-sim"), "log", "-1",
                                                        "--format=%h")), "", ""]]:
        for j, x in enumerate(p_):
            ws.cell(row=row, column=j + 1, value=x)
        row += 1
    # layout
    widths = [26, 22, 60, 60, 60, 30] + [14] * (len(COLS) - 6)
    for j, wdt in enumerate(widths):
        ws.column_dimensions[get_column_letter(j + 1)].width = wdt
    ws.freeze_panes = "E2"
    out = os.path.join(DOCS, "results", "Snake_all_versions_%s.xlsx" % date)
    wb.save(out)
    print(out)

    # ---------------------------------------------------------- verification
    wb2 = load_workbook(out, data_only=True)
    ws2 = wb2["Snake_results"]
    hdr = [c.value for c in ws2[h2]]
    rows2 = [[c.value for c in r_] for r_ in ws2.iter_rows(min_row=h2 + 1, max_row=last2)]
    get = lambda ver, gpu, cap, wl, col: next((r_[hdr.index(col)] for r_ in rows2 if r_[0] == ver and r_[1] == gpu
                                               and r_[2] == cap and str(r_[3]).startswith(wl)), None)
    G7 = "GMEAN / mean: our 7 apps"
    BK7 = "GMEAN best-kernel ratio (Diagnostic): our 7 apps"
    checks = [
        ("Tier 1 Snake whole-run GMEAN, 7 apps", -0.29, get("Tier 1 Snake", "PAPER_V100", "1B", G7,
                                                            "IPC gain vs no prefetching (%)"), 2),
        ("SK_R0_D0 PAPER_V100 1B whole-run GMEAN, 7 apps", -0.52, get("SK_R0_D0", "PAPER_V100", "1B", G7,
                                                                       "IPC gain vs no prefetching (%)"), 2),
        ("Best-kernel GMEAN, 7 apps, Tier 1 Snake", 2.05, (lambda v: (v - 1) * 100 if v else None)(
            get("Tier 1 Snake", "PAPER_V100", "1B", BK7, "Best kernel IPC / same kernel without prefetching (Diagnostic)")), 2),
        ("Best-kernel GMEAN, 7 apps, SK_R0_D0", 5.72, (lambda v: (v - 1) * 100 if v else None)(
            get("SK_R0_D0", "PAPER_V100", "1B", BK7, "Best kernel IPC / same kernel without prefetching (Diagnostic)")), 2),
        ("Tier 1 predicted coverage, 7 apps", 63.2, get("Tier 1 Snake", "PAPER_V100", "1B", G7,
                                                        "Coverage, Snake-paper style (coverage_pred, %)"), 1),
        ("Tier 1 timely accuracy, 7 apps", 6.1, get("Tier 1 Snake", "PAPER_V100", "1B", G7,
                                                    "Accuracy, Snake-paper style (timely / demand, %)"), 1),
        ("SK_SC base QV100, SnakeCharmer's 6 workloads, accuracy H/F", 66.2,
         get("SK_SC", "base QV100", "100M", "GMEAN / mean: SnakeCharmer's workloads (6, with bfs",
             "SnakeCharmer accuracy H/F (%)"), 1),
        ("histo best kernel ratio, Tier 1 Snake", 1.1148, get("Tier 1 Snake", "PAPER_V100", "1B", "histo",
                                                              "Best kernel IPC / same kernel without prefetching (Diagnostic)"), 4),
        ("histo best kernel ratio, SK_R0_D0", 1.4217, get("SK_R0_D0", "PAPER_V100", "1B", "histo",
                                                          "Best kernel IPC / same kernel without prefetching (Diagnostic)"), 4),
        ("histo best kernel ratio, B", 1.4332, get("B", "PAPER_V100", "1B", "histo",
                                                   "Best kernel IPC / same kernel without prefetching (Diagnostic)"), 4),
    ]
    fail = bool(bad)
    print("\nVerification (reopened workbook):")
    for name, exp, got, dec in checks:
        ok = isinstance(got, (int, float)) and round(got, dec) == round(exp, dec)
        fail |= not ok
        print("  %-60s expected %s got %s  %s" % (name, exp, ("%.*f" % (dec, got)) if isinstance(got, (int, float))
                                                  else got, "OK" if ok else "MISMATCH"))
    for b_ in bad:
        print("  CHECK FAILED:", b_)
    print("rows: block 1 %d, block 2 %d (header row %d), mismatch notes %d" % (len(vdef), last2 - h2, h2, len(mism)))
    sys.exit(1 if fail else 0)


if __name__ == "__main__":
    main()
